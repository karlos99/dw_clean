"""
State (TOMS) upload export for DESI.

Builds the CAASPP and ELPAC student-accommodation upload workbooks using the
exact TOMS template headers. Each worksheet has an SSID column plus one column
per designated support (by its upload column letter); a student's cell contains
the option's reporting code.

Sources merged per student:
* DESI-entered options (``StudentAccommodation``) for every support tagged
  CAASPP/ELPAC, regardless of whether the assessment is active.
* SEIS settings (``desi_seis``) for the matching program, matched by option code
  (SEIS wins when both provide a code for the same support).
"""

import io
import zipfile
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from .models import Accommodation, AccommodationOption, DesiSeis, StudentAccommodation
from .state_upload_templates import (
    CAASPP_SHEET_NAME,
    ELPAC_SHEET_NAME,
)

PROGRAM_TAGS = ('CAASPP', 'ELPAC')
TAG_TO_TEST_GROUP = {'CAASPP': 'caaspp', 'ELPAC': 'elpac'}
TEMPLATE_DIR = Path(__file__).resolve().parent / 'templates' / 'state_upload'
TEMPLATE_FILES = {
    'CAASPP': TEMPLATE_DIR / 'CAASPP_Upload_Stu_Accom_Template.xlsx',
    'ELPAC': TEMPLATE_DIR / 'ELPAC_Upload_Stu_Accom_Template.xlsx',
}


def _excel_letters(count):
    """Return the first ``count`` Excel column letters: A, B, …, Z, AA, …"""
    letters = []
    for i in range(1, count + 1):
        text = ''
        n = i
        while n:
            n, rem = divmod(n - 1, 26)
            text = chr(65 + rem) + text
        letters.append(text)
    return letters


def _letter_index(letter):
    """0-based index for an Excel column letter ('A'->0, 'AA'->26)."""
    letter = (letter or '').strip().upper()
    if not letter:
        return None
    value = 0
    for ch in letter:
        if not ch.isalpha():
            return None
        value = value * 26 + (ord(ch) - 64)
    return value - 1


def _supports_for_tag(tag):
    """All supports tagged ``tag`` ordered by their upload column letter."""
    supports = list(Accommodation.objects.filter(tag=tag))
    supports.sort(key=lambda s: (_letter_index(s.column_letter) is None,
                                 _letter_index(s.column_letter) or 0,
                                 s.name.casefold()))
    return supports


def _student_codes_for_tag(tag, students):
    """
    Map each student to ``{accommodation_id: option_code}`` for a program tag.

    DESI options are read in bulk; SEIS codes are matched by option code within
    the program tag and take precedence over DESI.
    """
    keys = [s.stu_annual_key for s in students]
    ssids = [s.ssid for s in students if s.ssid]

    desi_by_key = {}
    if keys:
        records = (
            StudentAccommodation.objects.filter(student_annual_key__in=keys)
            .prefetch_related('options__accommodation')
        )
        for record in records:
            bucket = desi_by_key.setdefault(record.student_annual_key, {})
            for option in record.options.all():
                accommodation = option.accommodation
                if accommodation and accommodation.tag == tag and option.code:
                    bucket[accommodation.id] = option.code

    seis_codes_by_ssid = {}
    group = TAG_TO_TEST_GROUP.get(tag)
    if group and ssids:
        try:
            with transaction.atomic():
                for ssid, settings in (
                    DesiSeis.objects.filter(ssid__in=ssids, test_group=group)
                    .values_list('ssid', 'settings')
                ):
                    codes = {
                        str(value).strip()
                        for value in (settings or {}).values()
                        if value and str(value).strip()
                    }
                    if codes:
                        seis_codes_by_ssid.setdefault(ssid, set()).update(codes)
        except Exception:
            seis_codes_by_ssid = {}

    option_by_code = {}
    all_codes = set()
    for codes in seis_codes_by_ssid.values():
        all_codes.update(codes)
    if all_codes:
        for option in (
            AccommodationOption.objects
            .filter(code__in=all_codes, accommodation__tag=tag)
            .select_related('accommodation')
        ):
            option_by_code.setdefault(option.code, option)

    results = []
    for student in students:
        codes = dict(desi_by_key.get(student.stu_annual_key, {}))
        if student.ssid and seis_codes_by_ssid.get(student.ssid):
            for code in seis_codes_by_ssid[student.ssid]:
                option = option_by_code.get(code)
                if option and option.accommodation_id:
                    codes[option.accommodation_id] = code
        results.append((student, codes))
    return results


def build_state_upload_workbook(tag, students):
    """Return the supplied TOMS template populated with the given students.

    The template is deliberately loaded instead of recreating a workbook.  The
    files contain instructions, specifications, definitions, formatting, and
    data validation that must remain intact in the download.  Only the named
    ``* Test Settings Template`` worksheet is populated.
    """
    if tag == 'CAASPP':
        sheet_name = CAASPP_SHEET_NAME
    elif tag == 'ELPAC':
        sheet_name = ELPAC_SHEET_NAME
    else:
        raise ValueError(f"Unknown program tag: {tag}")

    template_path = TEMPLATE_FILES[tag]
    if not template_path.is_file():
        raise FileNotFoundError(f"Missing state upload template: {template_path}")

    workbook = load_workbook(template_path)
    try:
        sheet = workbook[sheet_name]
    except KeyError as exc:
        raise ValueError(
            f"State upload template {template_path.name} has no {sheet_name!r} worksheet"
        ) from exc

    # Keep the template's header cells exactly as supplied.  In particular,
    # do not append columns for supports that are not part of the template.
    template_column_count = sheet.max_column
    letters = _excel_letters(template_column_count)
    letter_to_index = {letter: idx for idx, letter in enumerate(letters)}

    column_by_index = {}
    for support in _supports_for_tag(tag):
        idx = letter_to_index.get((support.column_letter or '').strip().upper())
        if idx is not None:
            column_by_index[idx] = support

    # Remove any sample/data values while preserving the template rows,
    # formatting, tables, comments, and validation rules.
    for row_number in range(2, sheet.max_row + 1):
        for column_number in range(1, template_column_count + 1):
            sheet.cell(row=row_number, column=column_number).value = None

    next_row = 2
    for student, codes in _student_codes_for_tag(tag, students):
        row_number = next_row
        next_row += 1
        row = [''] * template_column_count
        row[0] = student.ssid or ''
        for idx, support in column_by_index.items():
            row[idx] = codes.get(support.id, '')
        for column_number, value in enumerate(row, start=1):
            sheet.cell(row=row_number, column=column_number, value=value)

    # Make the data sheet the initially selected sheet, while retaining all
    # instructional/specification worksheets from the original template.
    workbook.active = workbook.index(sheet)

    return workbook


def _workbook_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_state_upload_zip(school_year, students, programs=PROGRAM_TAGS):
    """Return a ZIP file (bytes) containing one workbook per program tag."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        for tag in programs:
            data = _workbook_bytes(build_state_upload_workbook(tag, students))
            archive.writestr(f'{tag}_Upload_{school_year}.xlsx', data)
    return buffer.getvalue()
