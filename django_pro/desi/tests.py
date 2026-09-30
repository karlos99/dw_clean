from django.test import TestCase
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import connection
from django.urls import reverse
from django.core import signing
import csv
import io
from unittest import mock

from studata.models import CurrentStudent, Result, School as StudataSchool, Test
from users.models import Role
from desi.models import (
    Accommodation,
    AccommodationCategory,
    AccommodationOption,
    AccommodationResource,
    Assessment,
    AssessmentCriteria,
    DesiExportConfiguration,
    DesiSchoolAccess,
    DesiSeis,
    DesiSeisDisplaySettings,
    StudentAccommodation,
    StudentAssessmentCompletion,
    SupportRule,
    SupportRuleCriteria,
    SupportDefaultRule,
    SupportDefaultRuleCriteria,
)
from desi.views import (
    check_criteria,
    can_edit_desi,
    hidden_support_ids,
    default_support_locks,
    EXPORT_ENDPOINT_SIGNING_SALT,
)
from desi import importers

User = get_user_model()


def ensure_desi_seis_table():
    """Create the externally-managed ``desi_seis`` table for tests.

    ``DesiSeis`` is ``managed = False`` (the table is populated outside
    Django), so the test database does not create it. Tests that need SEIS
    rows create it here; the surrounding ``TestCase`` transaction rolls it
    back afterwards.
    """
    json_type = "jsonb" if connection.vendor == "postgresql" else "text"
    with connection.cursor() as cursor:
        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS desi_seis (
                id varchar(64) PRIMARY KEY,
                ssid varchar(32) NOT NULL,
                test_group varchar(32) NOT NULL,
                settings {json_type} NOT NULL,
                assessments {json_type} NOT NULL
            )
            """
        )


class DesiModelsAndEligibilityTests(TestCase):

    def setUp(self):
        # Create standard accommodation
        self.acc_tts = Accommodation.objects.create(
            name="Text-to-Speech",
            code="TTS",
            category="Presentation",
            question_text="Does the student need Text-to-Speech?"
        )
        self.opt_tts = AccommodationOption.objects.create(
            accommodation=self.acc_tts,
            name="Standard",
            code="TTS"
        )

        self.acc_et = Accommodation.objects.create(
            name="Extra Time 1.5x",
            code="ET1.5",
            category="Timing"
        )
        self.opt_et = AccommodationOption.objects.create(
            accommodation=self.acc_et,
            name="Standard",
            code="1.5X"
        )

        # Create assessment
        self.assessment = Assessment.objects.create(
            name="Smarter Balanced Math",
            subject="Math",
            is_active=True
        )
        self.assessment.supported_options.add(self.opt_tts, self.opt_et)

        # Setup criteria: SPED == 'Y' AND Grade in '03,04,05'
        self.criterion_sped = AssessmentCriteria.objects.create(
            assessment=self.assessment,
            field_name="sped",
            operator="is",
            value="Y"
        )
        self.criterion_grade = AssessmentCriteria.objects.create(
            assessment=self.assessment,
            field_name="grade",
            operator="in",
            value="03,04,05"
        )

        # Create students
        self.student_eligible = CurrentStudent.objects.create(
            stu_annual_key="eligible_key",
            ssid="ssid123",
            student_id="local123",
            student_name="Eligible Student",
            grade="03",
            sped="Y",
            school_year="2025-2026"
        )

        self.student_ineligible_grade = CurrentStudent.objects.create(
            stu_annual_key="ineligible_grade_key",
            ssid="ssid456",
            student_id="local456",
            student_name="Ineligible Grade",
            grade="06",
            sped="Y",
            school_year="2025-2026"
        )

        self.student_ineligible_sped = CurrentStudent.objects.create(
            stu_annual_key="ineligible_sped_key",
            ssid="ssid789",
            student_id="local789",
            student_name="Ineligible SPED",
            grade="03",
            sped="N",
            school_year="2025-2026"
        )

    def test_dynamic_criteria_evaluation(self):
        # 1. Eligible student tests
        eligible, failed = check_criteria(self.assessment.criteria.all(), self.student_eligible)
        self.assertTrue(eligible)
        self.assertEqual(len(failed), 0)

        # 2. Ineligible grade tests
        eligible, failed = check_criteria(self.assessment.criteria.all(), self.student_ineligible_grade)
        self.assertFalse(eligible)
        self.assertEqual(len(failed), 1)
        self.assertIn("Grade Level", failed[0])

        # 3. Ineligible SPED tests
        eligible, failed = check_criteria(self.assessment.criteria.all(), self.student_ineligible_sped)
        self.assertFalse(eligible)
        self.assertEqual(len(failed), 1)
        self.assertIn("Special Education", failed[0])


class DesiWorkflowAndViewsTests(TestCase):

    def setUp(self):
        # Roles setup
        self.admin_role = Role.objects.create(name="Project Admin")
        self.teacher_role = Role.objects.create(name="Teacher")

        # Users setup
        self.admin_user = User.objects.create_user(
            email="admin@test.com",
            username="admin@test.com",
            password="password123"
        )
        self.admin_user.roles.add(self.admin_role)

        self.teacher_user = User.objects.create_user(
            email="teacher@test.com",
            username="teacher@test.com",
            password="password123"
        )
        self.teacher_user.roles.add(self.teacher_role)

        # Assessment setup
        self.assessment_math = Assessment.objects.create(
            name="SBAC Math",
            is_active=True
        )
        self.assessment_ela = Assessment.objects.create(
            name="SBAC ELA",
            is_active=True
        )

        # Accommodation setup
        self.acc = Accommodation.objects.create(
            name="Text-to-Speech", 
            code="TTS",
            question_text="Does the student need Text-to-Speech?"
        )
        self.opt = AccommodationOption.objects.create(accommodation=self.acc, name="Standard", code="TTS")
        self.assessment_math.supported_options.add(self.opt)
        self.assessment_ela.supported_options.add(self.opt)

        # Student setup
        self.student = CurrentStudent.objects.create(
            stu_annual_key="student_key",
            ssid="ssid111",
            student_id="local111",
            student_name="Test Student",
            grade="04",
            school_year="2025-2026"
        )

    def test_accommodation_creation_support_mapping(self):
        self.client.force_login(self.admin_user)
        create_url = reverse('desi:create_accommodation')
        
        # Create Accommodation and check ELA and Math assessments with custom options
        post_data = {
            'name': 'Calculator',
            'code': 'CALC',
            'category': 'Response',
            'description': 'Graphing/Scientific calculator helper.',
            'question_text': 'Does the student require a calculator?',
            'bulk_options': 'Scientific, Graphing',
            'supported_assessments': [self.assessment_math.id, self.assessment_ela.id],
            f'assessment_opts_custom_{self.assessment_math.id}': 'Scientific Only, Basic 4-Function',
            f'assessment_opts_custom_{self.assessment_ela.id}': 'Graphing Only'
        }
        
        response = self.client.post(create_url, post_data)
        self.assertEqual(response.status_code, 302)

        # Verify options are created (2 master + 3 custom = 5 unique options)
        calc = Accommodation.objects.get(name='Calculator')
        self.assertEqual(calc.options.count(), 5)

        # Verify Math supports the correct custom options
        opt_sci_only = AccommodationOption.objects.get(accommodation=calc, name='Scientific Only')
        opt_basic = AccommodationOption.objects.get(accommodation=calc, name='Basic 4-Function')
        opt_grp_only = AccommodationOption.objects.get(accommodation=calc, name='Graphing Only')

        self.assertIn(opt_sci_only, self.assessment_math.supported_options.all())
        self.assertIn(opt_basic, self.assessment_math.supported_options.all())
        self.assertNotIn(opt_grp_only, self.assessment_math.supported_options.all())

        self.assertIn(opt_grp_only, self.assessment_ela.supported_options.all())
        self.assertNotIn(opt_sci_only, self.assessment_ela.supported_options.all())

    def test_assign_accommodations_post_with_assessment_completion(self):
        self.client.force_login(self.admin_user)
        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])
        
        # Check Standard ELA option and mark Math as completed but ELA as not completed
        post_data = {
            'student_options': [self.opt.id],
            'completed_assessments': [self.assessment_math.id]
        }
        
        response = self.client.post(url, post_data)
        self.assertEqual(response.status_code, 302)

        # Verify global option is saved
        student_acc = StudentAccommodation.objects.get(student_annual_key=self.student.stu_annual_key)
        self.assertIn(self.opt, student_acc.options.all())

        # Verify completion records
        comp_math = StudentAssessmentCompletion.objects.get(student_annual_key=self.student.stu_annual_key, assessment=self.assessment_math)
        self.assertTrue(comp_math.is_completed)

        comp_ela = StudentAssessmentCompletion.objects.get(student_annual_key=self.student.stu_annual_key, assessment=self.assessment_ela)
        self.assertFalse(comp_ela.is_completed)

    def test_assign_accommodations_htmx_post_returns_updated_sidebar_row(self):
        self.client.force_login(self.admin_user)
        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])

        post_data = {
            'student_options': [self.opt.id],
            'completed_assessments': [self.assessment_math.id, self.assessment_ela.id]
        }

        response = self.client.post(url, post_data, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="student-list-item-student_key"')
        self.assertContains(response, 'hx-swap-oob="true"')
        html = response.content.decode()
        self.assertIn('SBAC ELA', html)
        self.assertIn('SBAC Math', html)
        self.assertEqual(html.count('Completed</span>'), 2)

    def test_workspace_matrix_groups_and_merges_shared_supports(self):
        self.client.force_login(self.admin_user)
        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])

        # TTS has the same option 'Standard' on both tests, so it must appear as
        # a single row whose cell is merged across both test columns.
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

        columns = response.context['matrix_columns']
        self.assertEqual(len(columns), 2)
        self.assertEqual(
            {col['assessment_id'] for col in columns},
            {self.assessment_math.id, self.assessment_ela.id},
        )

        rows = [row for group in response.context['category_groups'] for row in group['rows']]
        tts_rows = [row for row in rows if row['accommodation'] == self.acc]
        self.assertEqual(len(tts_rows), 1)

        cells = [seg['cell'] for seg in tts_rows[0]['segments'] if seg['kind'] == 'cell']
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0]['colspan'], 2)
        self.assertEqual(
            set(cells[0]['test_ids']),
            {self.assessment_math.id, self.assessment_ela.id},
        )

    def test_admin_dashboard_saves_hidden_from_student_list(self):
        self.client.force_login(self.admin_user)
        response = self.client.post(reverse('desi:admin_dashboard'), {
            'action': 'save_student_list_hidden_supports',
            'hidden_from_student_list': [str(self.acc.id)],
        })
        self.assertEqual(response.status_code, 302)
        self.acc.refresh_from_db()
        self.assertTrue(self.acc.hide_from_student_list)

        response = self.client.post(reverse('desi:admin_dashboard'), {
            'action': 'save_student_list_hidden_supports',
        })
        self.assertEqual(response.status_code, 302)
        self.acc.refresh_from_db()
        self.assertFalse(self.acc.hide_from_student_list)

    def test_hidden_from_student_list_omits_workspace_but_not_export(self):
        self.client.force_login(self.admin_user)
        self.acc.hide_from_student_list = True
        self.acc.save(update_fields=['hide_from_student_list'])

        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        rows = [row for group in response.context['category_groups'] for row in group['rows']]
        self.assertEqual(rows, [])

        export_url = reverse('desi:export_data')
        support_col_key = f'support__{self.acc.id}'
        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'assessment_ids': [str(self.assessment_math.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            f'include__{support_col_key}': 'on',
            f'header__{support_col_key}': 'TTS Support',
            f'order__{support_col_key}': '2',
            'header__assessment_name': 'Assessment Name',
            'order__assessment_name': '99',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
        })
        self.assertEqual(response.status_code, 200)
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        self.assertEqual(rows[0], ['Student', 'TTS Support'])

    def test_create_accommodation_fails_when_no_assessments_exist(self):
        self.client.force_login(self.admin_user)
        
        # Delete all assessments
        Assessment.objects.all().delete()
        
        create_url = reverse('desi:create_accommodation')
        response = self.client.get(create_url)
        self.assertEqual(response.status_code, 302) # Redirect to dashboard
        
        # Verify the redirect goes to admin dashboard and shows error
        self.assertRedirects(response, reverse('desi:admin_dashboard'))

    def test_assign_accommodations_htmx_request(self):
        self.client.force_login(self.admin_user)
        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])

        # Get with HTMX headers
        response = self.client.get(url, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'desi/partials/workspace_form.html')
        self.assertTemplateNotUsed(response, 'desi/student_list.html')

    def test_edit_designated_support_option_name_and_code(self):
        self.client.force_login(self.admin_user)
        edit_url = reverse('desi:edit_accommodation', args=[self.acc.id])
        response = self.client.post(edit_url, {
            'name': self.acc.name,
            'code': self.acc.code,
            'category': self.acc.category,
            'description': self.acc.description,
            'video_url': '',
            'order': str(self.acc.order),
            'option_id[]': [str(self.opt.id)],
            'option_name[]': ['Read Aloud'],
            'option_code[]': ['RA'],
            'supported_assessments': [str(self.assessment_math.id), str(self.assessment_ela.id)],
            f'assessment_opts_{self.assessment_math.id}': ['Read Aloud'],
            f'assessment_opts_{self.assessment_ela.id}': ['Read Aloud'],
        })
        self.assertEqual(response.status_code, 302)
        self.opt.refresh_from_db()
        self.assertEqual(self.opt.name, 'Read Aloud')
        self.assertEqual(self.opt.code, 'RA')

    def test_edit_support_with_comma_option_names_can_save_resources(self):
        """Regression: option names may contain commas (e.g. '...Large, Black.').
        Editing such a support (here to add a resource) must not split the names
        or fail the unique-name validation."""
        self.client.force_login(self.admin_user)
        acc = Accommodation.objects.create(name="Mouse Pointer", code="MP")
        opts = []
        for name, code in [
            ("Mouse Pointer is selected as Large, Black.", "MP_L_BLK"),
            ("Mouse Pointer is selected as Large, White.", "MP_L_WHT"),
            ("Mouse Pointer is selected as Extra Large, Black.", "MP_XL_BLK"),
        ]:
            opts.append(AccommodationOption.objects.create(
                accommodation=acc, name=name, code=code
            ))
        for ast in (self.assessment_math, self.assessment_ela):
            ast.supported_options.add(*opts)

        edit_url = reverse('desi:edit_accommodation', args=[acc.id])
        response = self.client.post(edit_url, {
            'name': acc.name,
            'code': acc.code,
            'category': acc.category,
            'description': acc.description,
            'video_url': '',
            'order': str(acc.order),
            'option_id[]': [str(o.id) for o in opts],
            'option_name[]': [o.name for o in opts],
            'option_code[]': [o.code for o in opts],
            'supported_assessments': [str(self.assessment_math.id), str(self.assessment_ela.id)],
            f'assessment_opts_{self.assessment_math.id}': [o.name for o in opts],
            f'assessment_opts_{self.assessment_ela.id}': [o.name for o in opts],
            'resource_id[]': [''],
            'resource_label[]': ['Video guide'],
            'resource_url[]': ['https://www.youtube.com/watch?v=abc123'],
            'resource_kind[]': ['youtube'],
        })
        self.assertEqual(response.status_code, 302)
        acc.refresh_from_db()
        # The three options still exist intact (not split on commas).
        self.assertEqual(acc.options.count(), 3)
        self.assertTrue(acc.options.filter(name="Mouse Pointer is selected as Large, Black.").exists())
        self.assertTrue(acc.options.filter(name="Mouse Pointer is selected as Extra Large, Black.").exists())
        # Resource was saved.
        self.assertTrue(acc.resources.filter(url="https://www.youtube.com/watch?v=abc123").exists())

    def test_create_designated_support_with_dependency(self):
        self.client.force_login(self.admin_user)
        parent_support = Accommodation.objects.create(name="Read Aloud", code="RA")
        parent_option = AccommodationOption.objects.create(
            accommodation=parent_support,
            name="Student Reads Independently",
            code="NO_RA"
        )

        create_url = reverse('desi:create_accommodation')
        response = self.client.post(create_url, {
            'name': 'Human Reader',
            'code': 'HR',
            'category': 'Presentation',
            'description': 'Human reader support',
            'bulk_options': 'Standard',
            'depends_on_accommodation': str(parent_support.id),
            'depends_on_option': str(parent_option.id),
        })
        self.assertEqual(response.status_code, 302)

        created = Accommodation.objects.get(name='Human Reader')
        self.assertEqual(created.depends_on_accommodation_id, parent_support.id)
        self.assertEqual(created.depends_on_option_id, parent_option.id)

    def test_dependency_selection_hides_and_removes_saved_entry(self):
        self.client.force_login(self.admin_user)
        trigger_support = Accommodation.objects.create(name="Read Aloud", code="RA")
        trigger_option = AccommodationOption.objects.create(
            accommodation=trigger_support,
            name="Student Reads Independently",
            code="NO_RA"
        )
        dependent_support = Accommodation.objects.create(
            name="Human Reader",
            code="HR",
            depends_on_accommodation=trigger_support,
            depends_on_option=trigger_option,
        )
        dependent_option = AccommodationOption.objects.create(
            accommodation=dependent_support,
            name="Standard",
            code="HR_STD"
        )
        self.assessment_math.supported_options.add(trigger_option, dependent_option)

        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])
        response = self.client.post(url, {
            'student_options': [trigger_option.id, dependent_option.id],
        })
        self.assertEqual(response.status_code, 302)

        student_acc = StudentAccommodation.objects.get(student_annual_key=self.student.stu_annual_key)
        saved_option_ids = set(student_acc.options.values_list('id', flat=True))
        self.assertIn(trigger_option.id, saved_option_ids)
        self.assertNotIn(dependent_option.id, saved_option_ids)

    def test_dependency_scoped_to_some_assessments_does_not_remove_global_entry(self):
        self.client.force_login(self.admin_user)
        trigger_support = Accommodation.objects.create(name="Read Aloud Scoped", code="RAS")
        trigger_option = AccommodationOption.objects.create(
            accommodation=trigger_support,
            name="Independent",
            code="IND"
        )
        dependent_support = Accommodation.objects.create(
            name="Human Reader Scoped",
            code="HRS",
            depends_on_accommodation=trigger_support,
            depends_on_option=trigger_option,
        )
        dependent_support.dependency_assessments.add(self.assessment_math)
        dependent_option = AccommodationOption.objects.create(
            accommodation=dependent_support,
            name="Standard",
            code="STD"
        )
        self.assessment_math.supported_options.add(trigger_option, dependent_option)
        self.assessment_ela.supported_options.add(dependent_option)

        url = reverse('desi:assign_accommodations', args=[self.student.stu_annual_key])
        response = self.client.post(url, {
            'student_options': [trigger_option.id, dependent_option.id],
        })
        self.assertEqual(response.status_code, 302)

        student_acc = StudentAccommodation.objects.get(student_annual_key=self.student.stu_annual_key)
        saved_option_ids = set(student_acc.options.values_list('id', flat=True))
        self.assertIn(trigger_option.id, saved_option_ids)
        self.assertIn(dependent_option.id, saved_option_ids)

    def test_export_data_respects_column_rename_and_order(self):
        self.client.force_login(self.admin_user)
        export_url = reverse('desi:export_data')

        post_data = {
            'school_year': '2025-2026',
            'assessment_ids': [str(self.assessment_math.id)],
            'include__student_name': 'on',
            'header__student_name': 'Learner Name',
            'order__student_name': '2',
            'include__assessment_name': 'on',
            'header__assessment_name': 'Assessment',
            'order__assessment_name': '1',
        }
        # include defaults for non-selected columns
        for key, label, _group in [
            ('student_id', 'Local Student ID', 'student'),
            ('ssid', 'SSID', 'student'),
            ('stu_annual_key', 'Student Annual Key', 'student'),
            ('grade', 'Grade', 'student'),
            ('school_name', 'School', 'student'),
            ('school_year', 'School Year', 'student'),
            ('assessment_subject', 'Assessment Subject', 'support'),
            ('assessment_school_year', 'Assessment School Year', 'support'),
            ('assessment_completed', 'Assessment Completed', 'support'),
            ('support_name', 'Designated Support', 'support'),
            ('support_code', 'Support Code', 'support'),
            ('selected_option_name', 'Selected Option', 'support'),
            ('selected_option_code', 'Selected Option Code', 'support'),
            ('support_selected', 'Support Selected', 'support'),
        ]:
            post_data[f'header__{key}'] = label
            post_data[f'order__{key}'] = '99'

        response = self.client.post(export_url, post_data)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Type'].startswith('text/csv'))

        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)
        self.assertEqual(rows[0], ['Assessment', 'Learner Name'])
        self.assertEqual(rows[1][1], 'Test Student')

    def test_export_configuration_can_be_saved_and_reused(self):
        self.client.force_login(self.admin_user)
        export_url = reverse('desi:export_data')

        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'save_configuration': '1',
            'save_configuration_name': 'Math export',
            'save_configuration_description': 'Math assessment export',
            'assessment_ids': [str(self.assessment_math.id)],
            'support_value_field': 'code',
            'only_with_supports': 'on',
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            'header__assessment_name': 'Assessment',
            'order__assessment_name': '2',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
            'header__support_name': 'Designated Support',
            'order__support_name': '99',
            'header__support_code': 'Support Code',
            'order__support_code': '99',
            'header__selected_option_name': 'Selected Option',
            'order__selected_option_name': '99',
            'header__selected_option_code': 'Selected Option Code',
            'order__selected_option_code': '99',
            'header__support_selected': 'Support Selected',
            'order__support_selected': '99',
        })
        self.assertEqual(response.status_code, 200)
        config = DesiExportConfiguration.objects.get(owner=self.admin_user, name='Math export')
        self.assertEqual(config.config['school_year'], '2025-2026')
        self.assertEqual(config.config['support_value_field'], 'code')
        self.assertTrue(config.config['only_with_supports'])
        self.assertEqual(config.config['assessment_ids'], [self.assessment_math.id])

        load_response = self.client.get(export_url, {'config_id': config.id})
        self.assertEqual(load_response.status_code, 200)
        self.assertEqual(load_response.context['school_year'], '2025-2026')
        self.assertEqual(load_response.context['support_value_field'], 'code')
        self.assertTrue(load_response.context['only_with_supports'])
        self.assertEqual(load_response.context['selected_assessment_ids'], {self.assessment_math.id})

    def test_export_data_multiple_assessments_creates_multiple_rows(self):
        self.client.force_login(self.admin_user)
        export_url = reverse('desi:export_data')
        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'assessment_ids': [str(self.assessment_math.id), str(self.assessment_ela.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            'include__assessment_name': 'on',
            'header__assessment_name': 'Assessment',
            'order__assessment_name': '2',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
            'header__support_name': 'Designated Support',
            'order__support_name': '99',
            'header__support_code': 'Support Code',
            'order__support_code': '99',
            'header__selected_option_name': 'Selected Option',
            'order__selected_option_name': '99',
            'header__selected_option_code': 'Selected Option Code',
            'order__selected_option_code': '99',
            'header__support_selected': 'Support Selected',
            'order__support_selected': '99',
        })
        self.assertEqual(response.status_code, 200)
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        assessment_values = [r[1] for r in rows[1:]]
        self.assertIn('SBAC Math', assessment_values)
        self.assertIn('SBAC ELA', assessment_values)

    def test_export_data_supports_are_columns(self):
        self.client.force_login(self.admin_user)
        export_url = reverse('desi:export_data')
        support_col_key = f'support__{self.acc.id}'
        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'assessment_ids': [str(self.assessment_math.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            f'include__{support_col_key}': 'on',
            f'header__{support_col_key}': 'TTS Support',
            f'order__{support_col_key}': '2',
            'header__assessment_name': 'Assessment Name',
            'order__assessment_name': '99',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
        })
        self.assertEqual(response.status_code, 200)
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        self.assertEqual(rows[0], ['Student', 'TTS Support'])

    def test_export_data_support_columns_can_use_option_code(self):
        self.client.force_login(self.admin_user)
        StudentAccommodation.objects.create(student_annual_key=self.student.stu_annual_key).options.add(self.opt)
        export_url = reverse('desi:export_data')
        support_col_key = f'support__{self.acc.id}'
        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'support_value_field': 'code',
            'assessment_ids': [str(self.assessment_math.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            f'include__{support_col_key}': 'on',
            f'header__{support_col_key}': 'TTS Code',
            f'order__{support_col_key}': '2',
            'header__assessment_name': 'Assessment Name',
            'order__assessment_name': '99',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
        })
        self.assertEqual(response.status_code, 200)
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        self.assertEqual(rows[0], ['Student', 'TTS Code'])
        self.assertEqual(rows[1][1], 'TTS')

    def test_export_data_only_with_supports_skips_blank_rows(self):
        self.client.force_login(self.admin_user)
        StudentAccommodation.objects.create(student_annual_key=self.student.stu_annual_key).options.add(self.opt)
        CurrentStudent.objects.create(
            stu_annual_key="blank_student_key",
            ssid="blankssid",
            student_id="blankid",
            student_name="Blank Student",
            grade="04",
            school_year="2025-2026"
        )

        export_url = reverse('desi:export_data')
        response = self.client.post(export_url, {
            'school_year': '2025-2026',
            'only_with_supports': 'on',
            'assessment_ids': [str(self.assessment_math.id), str(self.assessment_ela.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            'include__assessment_name': 'on',
            'header__assessment_name': 'Assessment',
            'order__assessment_name': '2',
            'header__student_id': 'Local Student ID',
            'order__student_id': '99',
            'header__ssid': 'SSID',
            'order__ssid': '99',
            'header__stu_annual_key': 'Student Annual Key',
            'order__stu_annual_key': '99',
            'header__grade': 'Grade',
            'order__grade': '99',
            'header__school_name': 'School',
            'order__school_name': '99',
            'header__school_year': 'School Year',
            'order__school_year': '99',
            'header__assessment_subject': 'Assessment Subject',
            'order__assessment_subject': '99',
            'header__assessment_school_year': 'Assessment School Year',
            'order__assessment_school_year': '99',
            'header__assessment_completed': 'Assessment Completed',
            'order__assessment_completed': '99',
            f'header__support__{self.acc.id}': 'TTS',
            f'order__support__{self.acc.id}': '3',
        })
        self.assertEqual(response.status_code, 200)
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        exported_students = [r[0] for r in rows[1:]]
        self.assertIn('Test Student', exported_students)
        self.assertNotIn('Blank Student', exported_students)

    def test_export_endpoint_requires_valid_auth_token(self):
        endpoint_url = reverse('desi:export_data_endpoint')
        response = self.client.get(endpoint_url)
        self.assertEqual(response.status_code, 403)

    def test_export_endpoint_returns_csv_with_token_auth(self):
        endpoint_url = reverse('desi:export_data_endpoint')
        token = signing.dumps(
            {'uid': self.admin_user.id, 'purpose': 'desi_export'},
            salt=EXPORT_ENDPOINT_SIGNING_SALT
        )

        response = self.client.get(endpoint_url, {
            'auth': token,
            'school_year': '2025-2026',
            'assessment_ids': [str(self.assessment_math.id)],
            'include__student_name': 'on',
            'header__student_name': 'Student',
            'order__student_name': '1',
            'include__assessment_name': 'on',
            'header__assessment_name': 'Assessment',
            'order__assessment_name': '2',
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Type'].startswith('text/csv'))
        reader = csv.reader(io.StringIO(response.content.decode('utf-8')))
        rows = list(reader)
        self.assertEqual(rows[0], ['Student', 'Assessment'])

    def test_export_endpoint_uses_saved_config_id(self):
        config = DesiExportConfiguration.objects.create(
            owner=self.admin_user,
            name="Endpoint cfg",
            config={
                "school_year": "2025-2026",
                "assessment_ids": [self.assessment_math.id],
                "only_with_supports": False,
                "support_value_field": "code",
                "column_config": [
                    {"key": "student_name", "include": True, "header": "Student", "order": 1},
                    {"key": "assessment_name", "include": True, "header": "Assessment", "order": 2},
                ],
            },
        )
        token = signing.dumps(
            {"uid": self.admin_user.id, "purpose": "desi_export"},
            salt=EXPORT_ENDPOINT_SIGNING_SALT,
        )
        response = self.client.get(reverse("desi:export_data_endpoint"), {
            "auth": token,
            "config_id": str(config.id),
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Type"].startswith("text/csv"))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
        self.assertEqual(rows[0][0], "Student")
        self.assertIn("Assessment", rows[0])

    def test_export_page_endpoint_is_config_based(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("desi:export_data"))
        content = response.content.decode()
        # The endpoint URL references a saved config id rather than embedding the
        # whole form (which exceeded the server's request-line limit).
        self.assertIn("config_id", content)
        self.assertNotIn("new URLSearchParams(formData)", content)


CSV_SAMPLE = """Column,Display Name,Value,Value Definition,Content Area,Program,Setting Definition,Type
A,"Student SSID
(Example: 1234567890)",10 digits,NA,• all,•  all,This is a state-assigned student identifier.,NA
B,"Abacus (NON-EMBEDDED)
(Computer - M, S)",Blank (default),Abacus is not selected.,• math,"• Smarter math
• CAA math
• CAA science",Abacus resource text.,non-embedded accommodation
B,"Abacus (NON-EMBEDDED)
(Computer - M, S)",NEA_Abacus,Abacus is selected.,• math,"• Smarter math
• CAA math
• CAA science",Abacus resource text.,non-embedded accommodation
I,"*Color Contrast (EMBEDDED)
(Computer - E)",Blank (default),Color Contrast is not selected.,• ELA,• Smarter ELA,Color contrast text.,embedded designated support
I,"*Color Contrast (EMBEDDED)
(Computer - E)",TDS_CCInvert,Reverse Contrast,• ELA,• Smarter ELA,Color contrast text.,embedded designated support
I,"*Color Contrast (EMBEDDED)
(Computer - E)",TDS_CCMagenta,Black on Rose,• ELA,• Smarter ELA,Color contrast text.,embedded designated support
N,"*Noise Buffers (NON-EMBEDDED)
(Computer - M)",Blank (default),Noise Buffers is not selected.,• science,"• CAA Science
• CAST",Noise buffers text.,non-embedded designated support
N,"*Noise Buffers (NON-EMBEDDED)
(Computer - M)",NEDS_NoiseBuf,Noise Buffers is selected.,• science,"• CAA Science
• CAST",Noise buffers text.,non-embedded designated support
"""


class DesiCsvImporterTests(TestCase):

    def _parsed(self):
        return importers.parse_rows(csv.DictReader(io.StringIO(CSV_SAMPLE)))

    def test_parse_groups_supports_and_dedupes_assessments(self):
        parsed = self._parsed()

        # SSID (Type = NA) and "Blank (default)" rows are skipped.
        # 'CAA Science' (capital) on Noise Buffers folds into the earlier
        # 'CAA science' spelling rather than creating a new assessment.
        assessment_names = [a["name"] for a in parsed["assessments"]]
        self.assertEqual(
            assessment_names,
            ["Smarter math", "CAA math", "CAA science", "Smarter ELA", "CAST"],
        )

        self.assertEqual(len(parsed["supports"]), 3)
        by_letter = {s["column_letter"]: s for s in parsed["supports"]}
        self.assertEqual(by_letter["B"]["name"], "Abacus")
        self.assertEqual(by_letter["I"]["name"], "Color Contrast")
        self.assertEqual(by_letter["N"]["name"], "Noise Buffers")
        self.assertEqual(by_letter["N"]["category"], "non-embedded designated support")

        # "Blank (default)" rows do not become options
        self.assertEqual(len(by_letter["B"]["options"]), 1)
        self.assertEqual(by_letter["B"]["options"][0]["code"], "NEA_Abacus")
        self.assertEqual(len(by_letter["I"]["options"]), 2)
        self.assertEqual(
            [o["code"] for o in by_letter["I"]["options"]],
            ["TDS_CCInvert", "TDS_CCMagenta"],
        )

    def test_import_builds_assessments_supports_and_mappings(self):
        parsed = self._parsed()
        summary = importers.import_configuration(parsed, clear=True)

        self.assertEqual(summary, {"assessments": 5, "supports": 3, "options": 4})
        self.assertEqual(Assessment.objects.count(), 5)
        self.assertEqual(Accommodation.objects.count(), 3)
        self.assertEqual(AccommodationOption.objects.count(), 4)

        # Assessments are imported inactive by default
        self.assertFalse(Assessment.objects.filter(is_active=True).exists())

        caa_science = Assessment.objects.get(name="CAA science")
        cast = Assessment.objects.get(name="CAST")
        smarter_math = Assessment.objects.get(name="Smarter math")

        noise = AccommodationOption.objects.get(code="NEDS_NoiseBuf")
        self.assertIn(noise, caa_science.supported_options.all())
        self.assertIn(noise, cast.supported_options.all())
        self.assertNotIn(noise, smarter_math.supported_options.all())

        # The option is present on the assessments listed in its Program cell
        abacus = AccommodationOption.objects.get(code="NEA_Abacus")
        self.assertIn(abacus, smarter_math.supported_options.all())
        self.assertNotIn(abacus, cast.supported_options.all())

    def test_import_can_activate_assessments(self):
        parsed = self._parsed()
        importers.import_configuration(parsed, clear=True, active=True)
        self.assertEqual(Assessment.objects.filter(is_active=True).count(), 5)

    def test_import_is_rebuildable_without_duplicates(self):
        parsed = self._parsed()
        importers.import_configuration(parsed, clear=True)
        importers.import_configuration(parsed, clear=True)

        self.assertEqual(Assessment.objects.count(), 5)
        self.assertEqual(Accommodation.objects.count(), 3)
        self.assertEqual(AccommodationOption.objects.count(), 4)
        self.assertEqual(Accommodation.objects.filter(name="Color Contrast").count(), 1)

    def test_import_is_additive_and_stamps_tag(self):
        # Pre-existing manual records must survive an import (no wipe by default).
        Assessment.objects.create(name="Manual Test", subject="X", is_active=True)
        Accommodation.objects.create(name="Manual Support", code="M")

        parsed = self._parsed()
        importers.import_configuration(parsed, tag="CAASPP")

        self.assertTrue(Assessment.objects.filter(name="Manual Test").exists())
        self.assertTrue(Accommodation.objects.filter(name="Manual Support").exists())

        # Imported records carry the tag.
        self.assertEqual(Assessment.objects.get(name="Smarter math").tag, "CAASPP")
        self.assertEqual(Accommodation.objects.filter(tag="CAASPP").count(), 3)

        # Existing assessments are reused, not duplicated.
        self.assertEqual(Assessment.objects.count(), 6)

    def test_import_same_tag_is_idempotent(self):
        parsed = self._parsed()
        importers.import_configuration(parsed, tag="CAASPP")
        importers.import_configuration(parsed, tag="CAASPP")

        self.assertEqual(Assessment.objects.count(), 5)
        self.assertEqual(Accommodation.objects.count(), 3)
        self.assertEqual(AccommodationOption.objects.count(), 4)
        self.assertEqual(Accommodation.objects.filter(tag="CAASPP").count(), 3)

    def test_import_different_tags_keep_separate_supports(self):
        parsed = self._parsed()
        importers.import_configuration(parsed, tag="CAASPP")
        importers.import_configuration(parsed, tag="ELPAC")

        # Same assessment names are deduped across tags, but designated
        # supports are kept separate per source file.
        self.assertEqual(Assessment.objects.count(), 5)
        self.assertEqual(Accommodation.objects.count(), 6)
        self.assertEqual(AccommodationOption.objects.count(), 8)
        self.assertEqual(set(Accommodation.objects.values_list("tag", flat=True)), {"CAASPP", "ELPAC"})

    def test_map_only_adds_missing_associations_without_other_changes(self):
        parsed = self._parsed()

        # Pre-seed the tagged support + option + assessments.
        accommodation = Accommodation.objects.create(
            name="Abacus", code="B", column_letter="B", tag="CAASPP"
        )
        option = AccommodationOption.objects.create(
            accommodation=accommodation, name="Abacus is selected.", code="NEA_Abacus"
        )
        smarter = Assessment.objects.create(name="Smarter math", is_active=True)
        caa_math = Assessment.objects.create(name="CAA math", is_active=True)
        Assessment.objects.create(name="CAA science", is_active=True)
        unrelated = Assessment.objects.create(name="Unrelated Test", is_active=False)
        smarter.supported_options.add(option)  # already mapped

        summary = importers.update_assessment_mappings(parsed, tag="CAASPP", apply=True)

        # Added the two missing associations; existing one is untouched.
        self.assertEqual(summary["added_mappings"], 2)
        self.assertTrue(caa_math.supported_options.filter(id=option.id).exists())
        self.assertTrue(Assessment.objects.get(name="CAA science").supported_options.filter(id=option.id).exists())
        self.assertTrue(smarter.supported_options.filter(id=option.id).exists())
        self.assertFalse(unrelated.supported_options.filter(id=option.id).exists())

        # Nothing else was created or altered.
        self.assertEqual(Assessment.objects.count(), 4)
        self.assertEqual(Accommodation.objects.count(), 1)
        self.assertEqual(AccommodationOption.objects.count(), 1)

        # Idempotent: re-running adds nothing.
        summary2 = importers.update_assessment_mappings(parsed, tag="CAASPP", apply=True)
        self.assertEqual(summary2["added_mappings"], 0)

    def test_map_only_falls_back_when_tag_differs(self):
        parsed = self._parsed()

        # Support imported under the DEFAULT filename tag (e.g. CAASPP_Upload).
        accommodation = Accommodation.objects.create(
            name="Abacus", code="B", column_letter="B", tag="CAASPP_Upload"
        )
        option = AccommodationOption.objects.create(
            accommodation=accommodation, name="Abacus is selected.", code="NEA_Abacus"
        )
        for name in ("Smarter math", "CAA math", "CAA science"):
            Assessment.objects.create(name=name, is_active=True)

        # Running with a different --tag still adds the missing mappings.
        summary = importers.update_assessment_mappings(parsed, tag="CAASPP", apply=True)
        self.assertEqual(summary["added_mappings"], 3)
        for name in ("Smarter math", "CAA math", "CAA science"):
            self.assertTrue(Assessment.objects.get(name=name).supported_options.filter(id=option.id).exists())


class DesiTestingCoordinatorAndRulesTests(TestCase):

    def setUp(self):
        self.teacher_role = Role.objects.create(name="Teacher")
        self.coordinator_role = Role.objects.create(name="Testing Coordinator")
        self.project_role = Role.objects.create(name="Project Admin")

        self.user = User.objects.create_user(
            username="coord-teacher@test.com",
            email="coord-teacher@test.com",
            password="password123",
        )
        self.user.roles.add(self.teacher_role, self.coordinator_role)
        self.user.schools = ["SCH-A"]
        self.user.save()

        self.student = CurrentStudent.objects.create(
            stu_annual_key="desi_coord_student",
            ssid="coordssid",
            student_id="coordlocal",
            student_name="Coordinator Student",
            grade="03",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )

    def test_coordinator_addon_role_is_schoolwide_even_when_teacher_active(self):
        # Simulate middleware attaching an active role of Teacher.
        self.user.active_role = self.teacher_role
        self.assertIsNone(self.user.get_accessible_student_keys())

    def test_can_edit_matrix(self):
        coord = User.objects.create_user(username="coord@test.com", email="coord@test.com", password="password123")
        coord.roles.add(self.coordinator_role)
        coord.schools = ["SCH-A"]
        coord.save()
        self.assertTrue(can_edit_desi(coord, student=self.student))

        # Teacher default = edit.
        self.assertTrue(can_edit_desi(self.user, student=self.student))

        # View-only school flips teacher to read-only.
        DesiSchoolAccess.objects.create(school_state_id="SCH-A", teachers_can_edit=False)
        self.assertFalse(can_edit_desi(self.user, student=self.student))

    def test_view_only_school_blocks_teacher_save(self):
        DesiSchoolAccess.objects.create(school_state_id="SCH-A", teachers_can_edit=False)
        assessment = Assessment.objects.create(name="Blocked Test", is_active=True)
        self.user.active_role = self.teacher_role
        self.client.force_login(self.user)

        url = reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        response = self.client.post(url, {"completed_assessments": [str(assessment.id)]})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            StudentAssessmentCompletion.objects.filter(
                student_annual_key=self.student.stu_annual_key, assessment=assessment
            ).exists()
        )

    def test_support_rule_hides_support(self):
        accommodation = Accommodation.objects.create(name="Calculator", code="CALC")
        option = AccommodationOption.objects.create(accommodation=accommodation, name="Standard", code="STD")
        assessment = Assessment.objects.create(name="Hidden Math", subject="Math", is_active=True)
        assessment.supported_options.add(option)

        rule = SupportRule.objects.create(name="No grade 03 calculators", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="in", value="04,05")

        # Student is grade 03 -> criteria fail -> support hidden.
        hidden = hidden_support_ids(self.student)
        self.assertIn(accommodation.id, hidden)

        # Grade 04 student keeps the support.
        grade4 = CurrentStudent.objects.create(
            stu_annual_key="grade4",
            ssid="g4ssid",
            student_id="g4",
            student_name="Grade Four",
            grade="04",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(grade4))

    def test_support_rule_hides_support_in_workspace_matrix(self):
        project = User.objects.create_user(username="rule_view@test.com", email="rule_view@test.com", password="password123")
        project.roles.add(self.project_role)
        self.client.force_login(project)

        accommodation = Accommodation.objects.create(name="Calculator Matrix", code="CALCM")
        option = AccommodationOption.objects.create(accommodation=accommodation, name="Standard", code="STD")
        assessment = Assessment.objects.create(name="Matrix Math", subject="Math", is_active=True)
        assessment.supported_options.add(option)

        rule = SupportRule.objects.create(name="Grade 04-05 only", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="in", value="04,05")

        response = self.client.get(reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]))
        self.assertEqual(response.status_code, 200)
        rows = [
            row for group in response.context["category_groups"]
            for row in group["rows"]
        ]
        self.assertEqual(rows, [])

        grade5 = CurrentStudent.objects.create(
            stu_annual_key="grade5_matrix",
            ssid="g5ssid",
            student_id="g5",
            student_name="Grade Five",
            grade="05",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )
        response = self.client.get(reverse("desi:assign_accommodations", args=[grade5.stu_annual_key]))
        rows = [
            row for group in response.context["category_groups"]
            for row in group["rows"]
        ]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["accommodation"].id, accommodation.id)

    def test_support_rule_grade_normalizes_unpadded_values(self):
        accommodation = Accommodation.objects.create(name="GradeNorm", code="GN")
        rule = SupportRule.objects.create(name="Grade 4 list", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="in", value="04,05")

        stu = CurrentStudent.objects.create(
            stu_annual_key="gn_key",
            ssid="gn_ssid",
            student_id="gn",
            student_name="Grade Norm",
            grade="4",
            school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu))

    def test_inactive_support_rule_does_not_hide(self):
        accommodation = Accommodation.objects.create(name="InactiveRuleSupport", code="IRS")
        rule = SupportRule.objects.create(name="Inactive", is_active=False)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="is", value="99")
        self.assertNotIn(accommodation.id, hidden_support_ids(self.student))

    def test_support_rule_form_creates_rule(self):
        project = User.objects.create_user(username="p@test.com", email="p@test.com", password="password123")
        project.roles.add(self.project_role)
        self.client.force_login(project)
        accommodation = Accommodation.objects.create(name="Masking", code="M")

        response = self.client.post(reverse("desi:support_rule_create"), {
            "name": "Masking only grade 05",
            "description": "",
            "is_active": "on",
            "supports": [str(accommodation.id)],
            "crit_field[]": ["grade"],
            "crit_operator[]": ["is"],
            "crit_value[]": ["05"],
        })
        self.assertEqual(response.status_code, 302)
        rule = SupportRule.objects.get(name="Masking only grade 05")
        self.assertIn(accommodation, rule.supports.all())
        self.assertEqual(rule.criteria.count(), 1)

    def test_support_rule_with_multiple_and_criteria(self):
        accommodation = Accommodation.objects.create(name="SpeechToText", code="STT")
        rule = SupportRule.objects.create(name="Grade 04 SPED only", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="is", value="04", logic_operator="AND", order=0)
        SupportRuleCriteria.objects.create(rule=rule, field_name="sped", operator="is", value="Y", logic_operator="AND", order=1)

        # Student Grade 03, SPED N -> fails both
        self.assertIn(accommodation.id, hidden_support_ids(self.student))

        # Student Grade 04, SPED Y -> passes both
        stu_pass = CurrentStudent.objects.create(
            stu_annual_key="pass_and_key",
            ssid="pass_and_ssid",
            student_id="pass_and",
            student_name="Pass AND",
            grade="04",
            sped="Y",
            school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_pass))

        # Student Grade 04, SPED N -> fails second
        stu_fail_sped = CurrentStudent.objects.create(
            stu_annual_key="fail_sped_key",
            ssid="fail_sped_ssid",
            student_id="fail_sped",
            student_name="Fail SPED",
            grade="04",
            sped="N",
            school_year="2025-2026",
        )
        self.assertIn(accommodation.id, hidden_support_ids(stu_fail_sped))

    def test_support_rule_with_or_criteria(self):
        accommodation = Accommodation.objects.create(name="ReadAloud", code="RA")
        rule = SupportRule.objects.create(name="SPED or 504", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="sped", operator="is", value="Y", logic_operator="AND", order=0)
        SupportRuleCriteria.objects.create(rule=rule, field_name="in_504", operator="is", value="Y", logic_operator="OR", order=1)

        # Student with SPED=Y, 504=N -> eligible via first branch
        stu_sped = CurrentStudent.objects.create(
            stu_annual_key="or_sped_key",
            ssid="or_sped_ssid",
            student_id="or_sped",
            student_name="SPED Only",
            grade="03",
            sped="Y",
            in_504="N",
            school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_sped))

        # Student with SPED=N, 504=Y -> eligible via second branch
        stu_504 = CurrentStudent.objects.create(
            stu_annual_key="or_504_key",
            ssid="or_504_ssid",
            student_id="or_504",
            student_name="504 Only",
            grade="03",
            sped="N",
            in_504="Y",
            school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_504))

        # Student with SPED=N, 504=N -> fails both branches -> hidden
        stu_neither = CurrentStudent.objects.create(
            stu_annual_key="or_neither_key",
            ssid="or_neither_ssid",
            student_id="or_neither",
            student_name="Neither",
            grade="03",
            sped="N",
            in_504="N",
            school_year="2025-2026",
        )
        self.assertIn(accommodation.id, hidden_support_ids(stu_neither))

    def test_support_rule_with_mixed_and_or_criteria(self):
        # (Grade is 03 AND SPED is Y) OR (Grade is 04 AND in_504 is Y)
        accommodation = Accommodation.objects.create(name="Braille", code="BR")
        rule = SupportRule.objects.create(name="Compound Rule", is_active=True)
        rule.supports.add(accommodation)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="is", value="03", logic_operator="AND", order=0)
        SupportRuleCriteria.objects.create(rule=rule, field_name="sped", operator="is", value="Y", logic_operator="AND", order=1)
        SupportRuleCriteria.objects.create(rule=rule, field_name="grade", operator="is", value="04", logic_operator="OR", order=2)
        SupportRuleCriteria.objects.create(rule=rule, field_name="in_504", operator="is", value="Y", logic_operator="AND", order=3)

        # Grade 03 & SPED Y -> passes branch 1
        stu_b1 = CurrentStudent.objects.create(
            stu_annual_key="b1_key", ssid="b1_ssid", student_id="b1", student_name="Branch 1",
            grade="03", sped="Y", in_504="N", school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_b1))

        # Grade 04 & 504 Y -> passes branch 2
        stu_b2 = CurrentStudent.objects.create(
            stu_annual_key="b2_key", ssid="b2_ssid", student_id="b2", student_name="Branch 2",
            grade="04", sped="N", in_504="Y", school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_b2))

        # Grade 03 & 504 Y (but SPED N) -> fails branch 1 and branch 2
        stu_cross = CurrentStudent.objects.create(
            stu_annual_key="cross_key", ssid="cross_ssid", student_id="cross", student_name="Cross",
            grade="03", sped="N", in_504="Y", school_year="2025-2026",
        )
        self.assertIn(accommodation.id, hidden_support_ids(stu_cross))

    def test_support_rule_form_creates_and_edits_or_criteria(self):
        project = User.objects.create_user(username="proj_admin@test.com", email="proj_admin@test.com", password="password123")
        project.roles.add(self.project_role)
        self.client.force_login(project)
        accommodation = Accommodation.objects.create(name="Color Contrast", code="CC")

        # 1. Create with OR logic
        response = self.client.post(reverse("desi:support_rule_create"), {
            "name": "SPED OR 504 Rule",
            "description": "Accessible to either SPED or 504",
            "is_active": "on",
            "supports": [str(accommodation.id)],
            "crit_field[]": ["sped", "in_504"],
            "crit_operator[]": ["is", "is"],
            "crit_value[]": ["Y", "Y"],
            "crit_logic_op[]": ["AND", "OR"],
        })
        self.assertEqual(response.status_code, 302)
        rule = SupportRule.objects.get(name="SPED OR 504 Rule")
        criteria = list(rule.criteria.all())
        self.assertEqual(len(criteria), 2)
        self.assertEqual(criteria[0].logic_operator, "AND")
        self.assertEqual(criteria[0].order, 0)
        self.assertEqual(criteria[1].logic_operator, "OR")
        self.assertEqual(criteria[1].order, 1)
        self.assertIn("OR", rule.criteria_summary)

        # 2. Edit rule to change second condition to AND
        response_edit = self.client.post(reverse("desi:support_rule_edit", args=[rule.id]), {
            "name": "SPED AND 504 Rule",
            "description": "Now both must match",
            "is_active": "on",
            "supports": [str(accommodation.id)],
            "crit_field[]": ["sped", "in_504"],
            "crit_operator[]": ["is", "is"],
            "crit_value[]": ["Y", "Y"],
            "crit_logic_op[]": ["AND", "AND"],
        })
        self.assertEqual(response_edit.status_code, 302)
        rule.refresh_from_db()
        criteria_updated = list(rule.criteria.all())
        self.assertEqual(criteria_updated[1].logic_operator, "AND")
        self.assertIn("AND", rule.criteria_summary)
        self.assertNotIn("OR", rule.criteria_summary)

    def test_support_rule_form_creates_rule_with_list_in_checkboxes(self):
        project = User.objects.create_user(username="list_admin@test.com", email="list_admin@test.com", password="password123")
        project.roles.add(self.project_role)
        self.client.force_login(project)
        accommodation = Accommodation.objects.create(name="MultiListSupport", code="MLS")

        # Post with 'in' operator and comma-separated list (as generated by multi-select checkboxes)
        response = self.client.post(reverse("desi:support_rule_create"), {
            "name": "Grades 03,04,05 OR SPED",
            "description": "Multi-value list rule",
            "is_active": "on",
            "supports": [str(accommodation.id)],
            "crit_field[]": ["grade", "sped"],
            "crit_operator[]": ["in", "is"],
            "crit_value[]": ["03,04,05", "Y"],
            "crit_logic_op[]": ["AND", "OR"],
        })
        self.assertEqual(response.status_code, 302)
        rule = SupportRule.objects.get(name="Grades 03,04,05 OR SPED")
        criteria = list(rule.criteria.all())
        self.assertEqual(len(criteria), 2)
        self.assertEqual(criteria[0].operator, "in")
        self.assertEqual(criteria[0].value, "03,04,05")
        self.assertEqual(criteria[1].logic_operator, "OR")

        # Student Grade 04, SPED N -> matches branch 1
        stu_gr4 = CurrentStudent.objects.create(
            stu_annual_key="gr4_key", ssid="gr4_ssid", student_id="gr4", student_name="Grade 4 Student",
            grade="04", sped="N", school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_gr4))

        # Student Grade 06, SPED Y -> matches branch 2
        stu_sped_gr6 = CurrentStudent.objects.create(
            stu_annual_key="gr6_key", ssid="gr6_ssid", student_id="gr6", student_name="Grade 6 SPED",
            grade="06", sped="Y", school_year="2025-2026",
        )
        self.assertNotIn(accommodation.id, hidden_support_ids(stu_sped_gr6))

        # Student Grade 06, SPED N -> matches neither branch -> hidden
        stu_gr6_non_sped = CurrentStudent.objects.create(
            stu_annual_key="gr6_non_key", ssid="gr6_non_ssid", student_id="gr6_non", student_name="Grade 6 Non SPED",
            grade="06", sped="N", school_year="2025-2026",
        )
        self.assertIn(accommodation.id, hidden_support_ids(stu_gr6_non_sped))

    def test_admin_dashboard_saves_assessment_column_order(self):
        admin = User.objects.create_user(username="ast_admin@test.com", email="ast_admin@test.com", password="password123")
        admin.roles.add(self.project_role)
        self.client.force_login(admin)

        ast_sci = Assessment.objects.create(name="Science", is_active=True, order=10)
        ast_math = Assessment.objects.create(name="Math", is_active=True, order=20)
        ast_ela = Assessment.objects.create(name="ELA", is_active=True, order=30)

        # POST reorder: Math first (1), ELA second (2), Science third (3)
        response = self.client.post(reverse("desi:admin_dashboard"), {
            "action": "save_assessment_order",
            "assessment_order": [str(ast_math.id), str(ast_ela.id), str(ast_sci.id)],
        })
        self.assertEqual(response.status_code, 302)

        ast_math.refresh_from_db()
        ast_ela.refresh_from_db()
        ast_sci.refresh_from_db()
        self.assertEqual(ast_math.order, 1)
        self.assertEqual(ast_ela.order, 2)
        self.assertEqual(ast_sci.order, 3)

        # Verify ordering query returns them in this custom order
        ordered = list(Assessment.objects.filter(id__in=[ast_sci.id, ast_math.id, ast_ela.id]))
        self.assertEqual([a.id for a in ordered], [ast_math.id, ast_ela.id, ast_sci.id])

    def test_assign_accommodations_renders_matrix_columns_in_configured_order(self):
        admin = User.objects.create_user(username="ast_user@test.com", email="ast_user@test.com", password="password123")
        admin.roles.add(self.project_role)
        self.client.force_login(admin)

        ast_sci = Assessment.objects.create(name="Order Science", is_active=True, order=3)
        ast_math = Assessment.objects.create(name="Order Math", is_active=True, order=1)
        ast_ela = Assessment.objects.create(name="Order ELA", is_active=True, order=2)

        response = self.client.get(reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]))
        self.assertEqual(response.status_code, 200)

        matrix_columns = response.context["matrix_columns"]
        col_names = [c["assessment"].name for c in matrix_columns if c["assessment"].name.startswith("Order ")]
        self.assertEqual(col_names, ["Order Math", "Order ELA", "Order Science"])

        # Check template styling: Supports Matrix, Save Supports, and equal sizing table
        content = response.content.decode()
        self.assertIn("Supports Matrix", content)
        self.assertNotIn("Designated Supports by Section", content)
        self.assertIn("Save Supports", content)
        self.assertIn("desi-matrix-table", content)
        self.assertIn("desi-matrix-col-item", content)

    def test_admin_dashboard_uses_submenu_tabs(self):
        admin = User.objects.create_user(username="collapse_admin@test.com", email="collapse_admin@test.com", password="password123")
        admin.roles.add(self.project_role)
        self.client.force_login(admin)

        response = self.client.get(reverse("desi:admin_dashboard"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()

        # Submenu tabs with tab state (no collapsible sections anymore).
        self.assertIn("x-data=\"{ tab: 'assessments', showSidebar: true }\"", content)
        self.assertIn("showSidebar = !showSidebar", content)
        self.assertIn("tab='assessments'", content)
        self.assertIn("tab='order'", content)
        self.assertIn("tab='supports'", content)
        self.assertIn("tab='seis'", content)
        self.assertIn('x-show="tab===\'assessments\'"', content)
        self.assertIn('x-show="tab===\'order\'"', content)
        self.assertIn('x-show="tab===\'supports\'"', content)
        self.assertIn('x-show="tab===\'seis\'"', content)
        # Support pool uses a wider responsive grid on large layouts.
        self.assertIn("xl:grid-cols-4 2xl:grid-cols-5", content)

    def test_resource_crud_and_embed(self):
        accommodation = Accommodation.objects.create(name="TTS", code="TTS")
        resource = AccommodationResource.objects.create(
            accommodation=accommodation,
            label="Guide",
            url="https://www.youtube.com/watch?v=abc123",
            kind="youtube",
        )
        self.assertIn("youtube.com/embed/abc123", resource.get_embed_url())

    def test_page_resource_opens_externally_not_in_iframe(self):
        accommodation = Accommodation.objects.create(name="Directions", code="DIR")
        resource = AccommodationResource.objects.create(
            accommodation=accommodation,
            label="Translated directions",
            url="https://www.caaspp-elpac.org/resources/preparation/accessibility-resources/translated-test-directions#sb-ela",
            kind="page",
        )
        self.assertIsNone(resource.get_embed_url())
        self.assertTrue(resource.opens_in_new_tab())

    def test_school_report_counts(self):
        from desi.report_utils import build_school_report

        assessment = Assessment.objects.create(name="Report Math", subject="Math", is_active=True)
        CurrentStudent.objects.create(
            stu_annual_key="r1",
            ssid="r1ssid",
            student_id="r1",
            student_name="Report One",
            grade="03",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )
        other = CurrentStudent.objects.create(
            stu_annual_key="r2",
            ssid="r2ssid",
            student_id="r2",
            student_name="Report Two",
            grade="03",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )
        StudentAssessmentCompletion.objects.create(
            student_annual_key="r1",
            assessment=assessment,
            is_completed=True,
        )

        report = build_school_report("SCH-A", "2025-2026")
        summary = {row["assessment"].id: row for row in report["summary"]}
        self.assertEqual(summary[assessment.id]["eligible"], 3)
        self.assertEqual(summary[assessment.id]["completed"], 1)
        self.assertEqual(summary[assessment.id]["pending"], 2)
        self.assertIn(other, report["pending_by_assessment"][assessment.id])
        self.assertIn(self.student, report["pending_by_assessment"][assessment.id])

    def test_entered_supports_report(self):
        from unittest.mock import patch

        from desi.report_utils import build_entered_supports_report

        accommodation = Accommodation.objects.create(name="Calculator", code="CALC")
        option = AccommodationOption.objects.create(accommodation=accommodation, name="Scientific", code="SCI")
        updater = User.objects.create_user(username="up@test.com", email="up@test.com", password="password123")
        record, _ = StudentAccommodation.objects.get_or_create(
            student_annual_key=self.student.stu_annual_key,
        )
        record.options.add(option)
        record.updated_by = updater
        record.save()

        with patch(
            "desi.report_utils.seis_support_detail_for_student",
            return_value={
                "labels": ["Text-to-Speech: Read aloud"],
                "option_ids": {999},
                "accommodation_ids": {888},
            },
        ):
            report = build_entered_supports_report("SCH-A", "2025-2026")
        rows = {row["ssid"]: row for row in report["rows"]}
        self.assertIn("Calculator: Scientific", rows["coordssid"]["supports"])
        self.assertIn("Text-to-Speech: Read aloud (SEIS)", rows["coordssid"]["supports"])
        self.assertTrue(rows["coordssid"]["has_seis_supports"])
        self.assertEqual(rows["coordssid"]["from_seis"], "Yes")
        self.assertEqual(rows["coordssid"]["updated_by"], "up@test.com")

    def test_entered_supports_report_support_filter(self):
        from desi.report_utils import build_entered_supports_report

        accommodation = Accommodation.objects.create(name="Masking", code="MASK")
        option = AccommodationOption.objects.create(accommodation=accommodation, name="Covers", code="COV")
        record, _ = StudentAccommodation.objects.get_or_create(
            student_annual_key=self.student.stu_annual_key,
        )
        record.options.add(option)

        other = CurrentStudent.objects.create(
            stu_annual_key="other_support_student",
            ssid="otherssid",
            student_id="other",
            student_name="Other Student",
            grade="03",
            school_state_id="SCH-A",
            school_year="2025-2026",
        )
        StudentAccommodation.objects.get_or_create(student_annual_key=other.stu_annual_key)

        report = build_entered_supports_report(
            "SCH-A",
            "2025-2026",
            support_filter=f"o{option.id}",
        )
        ssids = {row["ssid"] for row in report["rows"]}
        self.assertEqual(ssids, {"coordssid"})

    def test_view_only_school_blocks_coordinator_save(self):
        DesiSchoolAccess.objects.create(school_state_id="SCH-A", teachers_can_edit=False)
        coord = User.objects.create_user(username="coord2@test.com", email="coord2@test.com", password="password123")
        coord.roles.add(self.coordinator_role)
        coord.schools = ["SCH-A"]
        coord.save()
        self.assertFalse(can_edit_desi(coord, student=self.student))

    def test_coordinator_dashboard_access_tab_renders(self):
        coord = User.objects.create_user(username="cd@test.com", email="cd@test.com", password="password123")
        coord.roles.add(self.coordinator_role)
        coord.schools = ["SCH-A"]
        coord.save()
        self.client.force_login(coord)

        response = self.client.get(reverse("desi:coordinator_dashboard"), HTTP_HOST="localhost")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Support Admin")
        self.assertContains(response, "Teacher Editing Access by School")


class DesiMappingPreservationTests(TestCase):

    def setUp(self):
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(username="map@test.com", email="map@test.com", password="password123")
        self.user.roles.add(role)

        self.acc = Accommodation.objects.create(name="Text-to-Speech Mapping", code="TTSM")
        self.opt = AccommodationOption.objects.create(accommodation=self.acc, name="Standard", code="STD")

        self.active = Assessment.objects.create(name="Active Mapping Test", is_active=True)
        self.inactive = Assessment.objects.create(name="Inactive Mapping Test", is_active=False)
        self.extra = Assessment.objects.create(name="Extra Mapping Test", is_active=True)
        self.active.supported_options.add(self.opt)
        self.inactive.supported_options.add(self.opt)

    def _edit_post(self, map_action=None, supported=None):
        data = {
            "name": self.acc.name,
            "code": self.acc.code,
            "category": self.acc.category,
            "description": "",
            "order": "0",
            "option_id[]": [str(self.opt.id)],
            "option_name[]": ["Standard"],
            "option_code[]": ["STD"],
            "supported_assessments": supported or [],
        }
        if map_action:
            data["map_action"] = map_action
        self.client.force_login(self.user)
        return self.client.post(reverse("desi:edit_accommodation", args=[self.acc.id]), data)

    def test_replace_save_preserves_inactive_assessment_mapping(self):
        response = self._edit_post(supported=[str(self.extra.id)])
        self.assertEqual(response.status_code, 302)

        # Active assessments not selected are removed; inactive stays mapped.
        self.assertFalse(self.active.supported_options.filter(id=self.opt.id).exists())
        self.assertTrue(self.inactive.supported_options.filter(id=self.opt.id).exists())
        self.assertTrue(self.extra.supported_options.filter(id=self.opt.id).exists())

    def test_merge_keeps_existing_mappings_and_adds_selected(self):
        response = self._edit_post(map_action="merge", supported=[str(self.extra.id)])
        self.assertEqual(response.status_code, 302)

        # Existing mappings (active + inactive) are untouched; extra added.
        self.assertTrue(self.active.supported_options.filter(id=self.opt.id).exists())
        self.assertTrue(self.inactive.supported_options.filter(id=self.opt.id).exists())
        self.assertTrue(self.extra.supported_options.filter(id=self.opt.id).exists())

    def test_option_removed_by_admin_is_deleted_even_if_inactive_only(self):
        self.client.force_login(self.user)

        # An option that only serves the inactive assessment.
        inactive_only = AccommodationOption.objects.create(
            accommodation=self.acc, name="Inactive Only", code="INO"
        )
        self.inactive.supported_options.add(inactive_only)

        # Edit posts ONLY the Standard option row (not Inactive Only), which
        # means the admin removed it and it must be deleted.
        response = self.client.post(reverse("desi:edit_accommodation", args=[self.acc.id]), {
            "name": self.acc.name,
            "code": self.acc.code,
            "category": self.acc.category,
            "description": "",
            "order": "0",
            "option_id[]": [str(self.opt.id)],
            "option_name[]": ["Standard"],
            "option_code[]": ["STD"],
            "supported_assessments": [str(self.active.id)],
        })
        self.assertEqual(response.status_code, 302)

        self.assertFalse(
            AccommodationOption.objects.filter(id=inactive_only.id).exists()
        )
        self.assertTrue(self.active.supported_options.filter(id=self.opt.id).exists())

    def test_project_admin_can_edit_option_text_and_code(self):
        self.client.force_login(self.user)

        response = self.client.post(reverse("desi:edit_accommodation", args=[self.acc.id]), {
            "name": self.acc.name,
            "code": self.acc.code,
            "category": self.acc.category,
            "description": "",
            "order": "0",
            "option_id[]": [str(self.opt.id)],
            "option_name[]": ["Read Aloud"],
            "option_code[]": ["RA"],
            "supported_assessments": [str(self.active.id), str(self.inactive.id)],
        })
        self.assertEqual(response.status_code, 302)
        self.opt.refresh_from_db()
        self.assertEqual(self.opt.name, "Read Aloud")
        self.assertEqual(self.opt.code, "RA")

    def test_project_admin_can_delete_an_option(self):
        self.client.force_login(self.user)

        removable = AccommodationOption.objects.create(
            accommodation=self.acc, name="Removable", code="REM"
        )
        self.active.supported_options.add(removable)
        self.inactive.supported_options.add(removable)

        # Post ONLY the Standard option (Removable omitted = deleted).
        response = self.client.post(reverse("desi:edit_accommodation", args=[self.acc.id]), {
            "name": self.acc.name,
            "code": self.acc.code,
            "category": self.acc.category,
            "description": "",
            "order": "0",
            "option_id[]": [str(self.opt.id)],
            "option_name[]": ["Standard"],
            "option_code[]": ["STD"],
            "supported_assessments": [str(self.active.id), str(self.inactive.id)],
        })
        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            AccommodationOption.objects.filter(id=removable.id).exists()
        )


class DesiSeisDisplaySettingsTests(TestCase):
    """Admin SEIS display selection + workspace gating."""

    def setUp(self):
        self.role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="seis-admin@test.com",
            email="seis-admin@test.com",
            password="password123",
        )
        self.user.roles.add(self.role)
        self.client.force_login(self.user)

        # Assessments: one active, one inactive, one unrelated.
        self.ast_math = Assessment.objects.create(name="Smarter math", is_active=True)
        self.ast_ela = Assessment.objects.create(name="Smarter ELA", is_active=False)
        self.ast_other = Assessment.objects.create(name="Unrelated Test", is_active=True)

        self.acc = Accommodation.objects.create(name="Text-to-Speech", code="TTS")
        self.opt = AccommodationOption.objects.create(
            accommodation=self.acc, name="Standard", code="TTS"
        )
        self.ast_math.supported_options.add(self.opt)
        self.ast_ela.supported_options.add(self.opt)
        self.ast_other.supported_options.add(self.opt)

        self.student = CurrentStudent.objects.create(
            stu_annual_key="seis_display_student",
            ssid="seisdisplay",
            student_id="seis-local",
            student_name="SEIS Display Student",
            grade="04",
            school_year="2025-2026",
        )

    def test_admin_dashboard_saves_selection(self):
        url = reverse("desi:admin_dashboard")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "SEIS Display Settings")

        # Save a selection of two assessments (one inactive).
        response = self.client.post(url, {
            "seis_display_assessments": [str(self.ast_math.id), str(self.ast_ela.id)],
        })
        self.assertEqual(response.status_code, 302)

        settings = DesiSeisDisplaySettings.objects.first()
        self.assertIsNotNone(settings)
        self.assertEqual(
            set(settings.assessments.values_list("id", flat=True)),
            {self.ast_math.id, self.ast_ela.id},
        )

    def test_dashboard_echoes_saved_selection(self):
        import re
        settings, _ = DesiSeisDisplaySettings.objects.get_or_create(pk=1)
        settings.assessments.set([self.ast_math.id])
        response = self.client.get(reverse("desi:admin_dashboard"))
        html = response.content.decode()

        def is_checked(value):
            # value appears before the optional whitespace-separated "checked"
            pattern = re.compile(r'name="seis_display_assessments" value="%s"\s*(checked)?' % value)
            m = pattern.search(html)
            return bool(m and m.group(1))

        # The selected math row is checked; ELA is not.
        self.assertTrue(is_checked(self.ast_math.id))
        self.assertFalse(is_checked(self.ast_ela.id))

class DesiSeisDisplaySettingsViewTests(TestCase):
    """SEIS admin display selection applied on the student workspace view."""

    def setUp(self):
        from unittest import mock
        cache.delete('landing:project_config')
        ensure_desi_seis_table()
        self._patchers = []
        self._fake_map = {}

        def _fake_build(ssid, assessments):
            return self._fake_map, {acc_id for _a, acc_id in self._fake_map}

        patcher_build = mock.patch.object(
            __import__("desi.views", fromlist=["build_seis_option_map"]),
            "build_seis_option_map",
            side_effect=_fake_build,
        )
        self._patchers.append(patcher_build)
        for p in self._patchers:
            p.start()
        self.addCleanup(self._stop_patchers)

        self.role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="seis-view-admin@test.com",
            email="seis-view-admin@test.com",
            password="password123",
        )
        self.user.roles.add(self.role)
        self.client.force_login(self.user)

        self.ast_math = Assessment.objects.create(name="Smarter math", is_active=True)
        self.ast_ela = Assessment.objects.create(name="Smarter ELA", is_active=False)

        self.acc = Accommodation.objects.create(name="Text-to-Speech", code="TTS")
        self.opt = AccommodationOption.objects.create(
            accommodation=self.acc, name="Standard", code="TTS"
        )
        self.ast_math.supported_options.add(self.opt)
        self.ast_ela.supported_options.add(self.opt)

        self.student = CurrentStudent.objects.create(
            stu_annual_key="seis_view_student",
            ssid="seisview",
            student_id="seis-view",
            student_name="SEIS View Student",
            grade="04",
            school_year="2025-2026",
        )
        # Make the superuser/PA an admin that can see the student.
        self.user.is_superuser = True
        self.user.save()

    def _stop_patchers(self):
        for p in self._patchers:
            try:
                p.stop()
            except Exception:
                pass

    def _create_seis_row(self, row_id, assessments, settings=None):
        return DesiSeis.objects.create(
            id=row_id,
            ssid=self.student.ssid,
            test_group="caaspp",
            settings=settings or {"TTS": "TTS"},
            assessments=assessments,
        )

    def test_seis_assigned_tests_show_locks_and_readonly(self):
        # The SEIS source assigns both math (active) and ELA (inactive).
        self._fake_map = {
            (self.ast_math.id, self.acc.id): self.opt,
            (self.ast_ela.id, self.acc.id): self.opt,
        }
        self._create_seis_row("seis-row-1", ["Smarter math", "Smarter ELA"])

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()

        # Both are read-only SEIS columns listing only their SEIS supports.
        self.assertIn("SEIS \u00b7 read only", html)
        self.assertIn("Smarter math", html)
        self.assertIn("Smarter ELA", html)
        self.assertIn("Locked from SEIS", html)

    def test_test_not_assigned_in_seis_is_hidden(self):
        self._fake_map = {
            (self.ast_math.id, self.acc.id): self.opt,
            (self.ast_ela.id, self.acc.id): self.opt,
        }
        # SEIS only assigns math -> ELA (inactive) must not appear, even though
        # the option map still carries an entry for it.
        self._create_seis_row("seis-row-2", ["Smarter math"])

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()

        self.assertIn("SEIS \u00b7 read only", html)
        self.assertIn("Smarter math", html)
        self.assertNotIn("Smarter ELA", html)

    def test_no_test_assignment_shows_regular_tests_but_read_only(self):
        # "No_Test" carries no per-test SEIS settings, so the regular tests are
        # shown — but the student is in the SEIS list, so the workspace is read-only.
        self._create_seis_row("seis-row-3", ["No_Test"])

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)

        columns = response.context["matrix_columns"]
        names = [c["assessment"].name for c in columns]
        math_col = next(c for c in columns if c["assessment"].name == "Smarter math")
        self.assertFalse(math_col["is_seis"])
        self.assertNotIn("Smarter ELA", names)

        # Present in SEIS -> read-only even though the regular tests are shown.
        self.assertTrue(response.context["seis_readonly"])
        self.assertFalse(response.context["can_edit"])
        self.assertIn("This student is listed in SEIS", response.content.decode())

    def test_seis_readonly_post_is_blocked(self):
        self._create_seis_row("seis-row-7", ["No_Test"])
        response = self.client.post(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]),
            {"completed_assessments": [str(self.ast_math.id)]},
        )
        self.assertEqual(response.status_code, 302)
        # Nothing is saved for a SEIS-managed student.
        self.assertFalse(
            StudentAssessmentCompletion.objects.filter(
                student_annual_key=self.student.stu_annual_key,
                assessment=self.ast_math,
                is_completed=True,
            ).exists()
        )

    def test_readonly_seis_workspace_preserves_existing_supports(self):
        # A fully SEIS-managed workspace renders no support controls, so a save
        # must not wipe the student's existing DESI selections.
        self._fake_map = {
            (self.ast_math.id, self.acc.id): self.opt,
        }
        self._create_seis_row("seis-row-5", ["Smarter math"])

        extra = Accommodation.objects.create(name="Extra Support", code="EX")
        extra_opt = AccommodationOption.objects.create(
            accommodation=extra, name="Extra Option", code="EXO"
        )
        record, _ = StudentAccommodation.objects.get_or_create(
            student_annual_key=self.student.stu_annual_key
        )
        record.options.add(extra_opt)
        record.selected_accommodations.add(extra)

        response = self.client.post(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]),
            {"action": "save_accommodations"},
        )
        self.assertEqual(response.status_code, 302)

        record.refresh_from_db()
        self.assertIn(extra_opt, record.options.all())
        self.assertIn(extra, record.selected_accommodations.all())

    def test_student_list_shows_only_seis_assigned_tests(self):
        # A plain active test would normally be listed in the sidebar.
        Assessment.objects.create(name="Plain Active", is_active=True)
        self._create_seis_row("seis-row-6", ["Smarter math"])

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)

        row = next(
            s for s in response.context["page_obj"]
            if s.stu_annual_key == self.student.stu_annual_key
        )
        names = [st["assessment"].name for st in row.assessment_statuses]
        self.assertEqual(names, ["Smarter math"])
        self.assertTrue(row.assessment_statuses[0]["is_seis"])
        self.assertNotIn("Plain Active", names)
        self.assertNotIn("Smarter ELA", names)

        # The sidebar shows a SEIS marker rather than a Completed/Pending status.
        content = response.content.decode()
        self.assertIn(">SEIS</span>", content)

    def test_seis_columns_follow_configured_assessment_order(self):
        # SEIS assigns an inactive test (ordered first) and an active one
        # (ordered last); they must render in the configured column order.
        self.ast_ela.order = 10
        self.ast_ela.save(update_fields=["order"])
        self.ast_math.order = 30
        self.ast_math.save(update_fields=["order"])

        self._fake_map = {
            (self.ast_math.id, self.acc.id): self.opt,
            (self.ast_ela.id, self.acc.id): self.opt,
        }
        self._create_seis_row("seis-row-4", ["Smarter math", "Smarter ELA"])

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)

        columns = response.context["matrix_columns"]
        names = [c["assessment"].name for c in columns]
        self.assertEqual(names, ["Smarter ELA", "Smarter math"])
        self.assertTrue(columns[0]["is_seis"])
        self.assertTrue(columns[1]["is_seis"])


class DesiSectionCategoryTests(TestCase):
    """Section categories coexist with the original free-text category."""

    def setUp(self):
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="section@test.com",
            email="section@test.com",
            password="password123",
        )
        self.user.roles.add(role)
        self.client.force_login(self.user)

        self.assessment = Assessment.objects.create(name="Section Test", is_active=True)
        self.student = CurrentStudent.objects.create(
            stu_annual_key="section_student",
            ssid="section",
            student_id="section-local",
            student_name="Section Student",
            grade="04",
            school_year="2025-2026",
        )

    def test_create_support_keeps_category_and_creates_section(self):
        response = self.client.post(reverse("desi:create_accommodation"), {
            "name": "Read Aloud",
            "code": "RA",
            "category": "Presentation",
            "new_section_category": "Accessibility",
            "bulk_options": "Standard",
        })
        self.assertEqual(response.status_code, 302)

        acc = Accommodation.objects.get(name="Read Aloud")
        self.assertEqual(acc.category, "Presentation")
        self.assertIsNotNone(acc.section_category_id)
        self.assertEqual(acc.section_category.name, "Accessibility")

    def test_workspace_matrix_groups_are_collapsible(self):
        cat = AccommodationCategory.objects.create(name="Presentation", order=1)
        acc = Accommodation.objects.create(name="Read Aloud", section_category=cat, order=1)
        opt = AccommodationOption.objects.create(accommodation=acc, name="Standard", code="STD")
        self.assessment.supported_options.add(opt)

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        # Each organization section has a collapsible rows <tbody> and a
        # separate, always-visible header so it can be expanded again.
        self.assertIn("desi-category-group", content)
        self.assertIn("data-category-toggle", content)
        self.assertIn("data-category-target=", content)
        self.assertIn("Presentation", content)
        # The toggle header must NOT live inside the collapsible tbody (that
        # would hide the control when a section collapses).
        self.assertNotRegex(
            content,
            r'<tbody class="desi-category-group"[^>]*>\s*<tr[^>]*data-category-toggle',
        )

    def test_section_picker_reuses_existing_section(self):
        existing = AccommodationCategory.objects.create(name="Timing", order=1)
        response = self.client.post(reverse("desi:create_accommodation"), {
            "name": "Extra Time",
            "code": "ET",
            "category": "Timing",
            "section_category_id": str(existing.id),
            "bulk_options": "Standard",
        })
        self.assertEqual(response.status_code, 302)
        acc = Accommodation.objects.get(name="Extra Time")
        self.assertEqual(acc.section_category_id, existing.id)

    def test_admin_ordering_persists_sections_and_supports(self):
        cat_a = AccommodationCategory.objects.create(name="Alpha", order=1)
        cat_b = AccommodationCategory.objects.create(name="Beta", order=2)
        acc1 = Accommodation.objects.create(name="One", section_category=cat_a, order=1)
        acc2 = Accommodation.objects.create(name="Two", section_category=cat_a, order=2)
        acc3 = Accommodation.objects.create(name="Three", section_category=cat_b, order=1)

        response = self.client.post(reverse("desi:save_desi_order"), {
            "category_order": [str(cat_b.id), str(cat_a.id)],
            f"support_order_{cat_a.id}": [str(acc2.id), str(acc1.id)],
            f"support_order_{cat_b.id}": [str(acc3.id)],
        })
        self.assertEqual(response.status_code, 302)

        cat_a.refresh_from_db(); cat_b.refresh_from_db()
        acc1.refresh_from_db(); acc2.refresh_from_db()
        self.assertEqual(cat_b.order, 1)
        self.assertEqual(cat_a.order, 2)
        self.assertEqual(acc2.order, 1)
        self.assertEqual(acc1.order, 2)

    def test_workspace_matrix_groups_follow_section_order(self):
        cat_a = AccommodationCategory.objects.create(name="Alpha", order=2)
        cat_b = AccommodationCategory.objects.create(name="Beta", order=1)
        for acc in (
            Accommodation.objects.create(name="One", section_category=cat_a, order=1),
            Accommodation.objects.create(name="Two", section_category=cat_b, order=1),
        ):
            opt = AccommodationOption.objects.create(accommodation=acc, name="Standard")
            self.assessment.supported_options.add(opt)

        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key]),
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        groups = response.context["category_groups"]
        self.assertEqual([g["category"] for g in groups], [cat_b, cat_a])

    def test_save_order_moves_support_between_sections(self):
        cat_a = AccommodationCategory.objects.create(name="Alpha", order=1)
        cat_b = AccommodationCategory.objects.create(name="Beta", order=2)
        acc = Accommodation.objects.create(name="Mover", section_category=cat_a, order=1)

        # Posting the support under Beta's list moves it into Beta.
        response = self.client.post(reverse("desi:save_desi_order"), {
            "category_order": [str(cat_a.id), str(cat_b.id)],
            f"support_order_{cat_b.id}": [str(acc.id)],
        })
        self.assertEqual(response.status_code, 302)
        acc.refresh_from_db()
        self.assertEqual(acc.section_category_id, cat_b.id)
        self.assertEqual(acc.order, 1)

        # Posting it under ``none`` clears its section.
        response = self.client.post(reverse("desi:save_desi_order"), {
            "category_order": [str(cat_a.id), str(cat_b.id)],
            "support_order_none": [str(acc.id)],
        })
        self.assertEqual(response.status_code, 302)
        acc.refresh_from_db()
        self.assertIsNone(acc.section_category_id)

    def test_admin_dashboard_links_to_organization_page(self):
        response = self.client.get(reverse("desi:admin_dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("desi:support_organization"))

    def test_support_organization_page_renders_and_creates_section(self):
        response = self.client.get(reverse("desi:support_organization"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Support Organization")
        self.assertContains(response, "Create Section")

        response = self.client.post(reverse("desi:support_organization"), {"name": "Accessibility"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(AccommodationCategory.objects.filter(name="Accessibility").exists())

    def test_support_organization_reuses_existing_section_name(self):
        AccommodationCategory.objects.create(name="Timing", order=1)
        self.client.post(reverse("desi:support_organization"), {"name": "timing"})
        self.assertEqual(AccommodationCategory.objects.filter(name__iexact="timing").count(), 1)

    def test_support_organization_lists_carry_section_keys(self):
        cat = AccommodationCategory.objects.create(name="Alpha", order=1)
        Accommodation.objects.create(name="One", section_category=cat, order=1)
        Accommodation.objects.create(name="Loose", order=1)

        response = self.client.get(reverse("desi:support_organization"))
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        # Each support list declares the section it represents and the page
        # rewrites hidden input names on drop so moves persist.
        self.assertIn(f'data-section-key="{cat.id}"', html)
        self.assertIn('data-section-key="none"', html)
        self.assertIn("syncSupportInputNames", html)

    def test_delete_section_category_moves_supports_to_uncategorized(self):
        cat = AccommodationCategory.objects.create(name="To Delete", order=1)
        acc1 = Accommodation.objects.create(name="Support 1", section_category=cat, order=1)
        acc2 = Accommodation.objects.create(name="Support 2", section_category=cat, order=2)

        response = self.client.post(reverse("desi:delete_section_category", args=[cat.id]))
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, reverse("desi:support_organization"))

        # Verify category is deleted
        self.assertFalse(AccommodationCategory.objects.filter(id=cat.id).exists())

        # Verify accommodations are moved to non-section (section_category is None)
        acc1.refresh_from_db()
        acc2.refresh_from_db()
        self.assertIsNone(acc1.section_category)
        self.assertIsNone(acc2.section_category)

    def test_delete_section_category_requires_admin(self):
        cat = AccommodationCategory.objects.create(name="Protected Section", order=1)
        self.client.logout()
        User = get_user_model()
        user = User.objects.create_user(username='regular_desi_user', password='password')
        self.client.force_login(user)

        response = self.client.post(reverse("desi:delete_section_category", args=[cat.id]))
        self.assertRedirects(response, reverse("landing:index"))
        self.assertTrue(AccommodationCategory.objects.filter(id=cat.id).exists())



class DesiStudentListSchoolFilterTests(TestCase):
    """The DESI student-list school filter must stay switchable after a selection."""

    def setUp(self):
        # The project config is cached in shared Redis; isolate this test from
        # a stale production value so the school year resolves from the test DB.
        cache.delete('landing:project_config')
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="school-filter@test.com",
            email="school-filter@test.com",
            password="password123",
        )
        self.user.roles.add(role)
        self.client.force_login(self.user)

        self.school_a = StudataSchool.objects.create(
            dw_schoolkey=1, state_id="SCH-A", short_name="Alpha", school_name="Alpha High"
        )
        self.school_b = StudataSchool.objects.create(
            dw_schoolkey=2, state_id="SCH-B", short_name="Beta", school_name="Beta High"
        )
        CurrentStudent.objects.create(
            stu_annual_key="sf-a", ssid="sf-a", student_id="sf-a",
            student_name="Alpha Student", grade="03",
            school_state_id="SCH-A", school_year="2025-2026",
        )
        CurrentStudent.objects.create(
            stu_annual_key="sf-b", ssid="sf-b", student_id="sf-b",
            student_name="Beta Student", grade="03",
            school_state_id="SCH-B", school_year="2025-2026",
        )

    def _get(self, school=None):
        url = reverse("desi:assign_accommodations", args=["sf-a"])
        params = {"school": school} if school else None
        return self.client.get(url, params)

    def test_selecting_a_school_keeps_the_dropdown_switchable(self):
        # Without a school filter the dropdown is rendered.
        response = self._get()
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["school_locked"])
        self.assertIsNone(response.context["global_active_school"])
        self.assertIn('id="school"', response.content.decode())

        # Applying a school filter keeps the dropdown (not a locked input) so
        # the user can switch to another school.
        response = self._get("SCH-A")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["school_locked"])
        self.assertEqual(response.context["school_filter"], "SCH-A")
        content = response.content.decode()
        self.assertIn('id="school"', content)
        self.assertIn('value="SCH-A" selected', content)
        self.assertIn('id="student-list-item-sf-a"', content)
        self.assertNotIn('id="student-list-item-sf-b"', content)

    def test_switching_from_one_school_to_another_works(self):
        first = self._get("SCH-A")
        self.assertIn('value="SCH-A" selected', first.content.decode())

        second = self._get("SCH-B")
        content = second.content.decode()
        self.assertIn('id="school"', content)
        self.assertIn('value="SCH-B" selected', content)
        self.assertIn('id="student-list-item-sf-b"', content)
        self.assertNotIn('id="student-list-item-sf-a"', content)

    def test_global_active_school_locks_the_filter(self):
        session = self.client.session
        session["active_school_id"] = "SCH-A"
        session.save()

        response = self._get()
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["school_locked"])
        content = response.content.decode()
        # Locked: a disabled input replaces the dropdown.
        self.assertNotIn('id="school"', content)
        self.assertIn('value="SCH-A"', content)

    def test_student_list_shows_504_badge(self):
        student = CurrentStudent.objects.get(stu_annual_key="sf-a")
        student.in_504 = "Y"
        student.save(update_fields=["in_504"])

        response = self._get("SCH-A")
        self.assertEqual(response.status_code, 200)
        self.assertIn("504: Y", response.content.decode())


class DesiStudentListPeriodFilterTests(TestCase):
    """Period filter for non-elementary schools, sourced from the schedule table."""

    def setUp(self):
        # Isolate from a stale shared project-config cache.
        cache.delete('landing:project_config')
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="period-filter@test.com",
            email="period-filter@test.com",
            password="password123",
        )
        self.user.roles.add(role)
        self.client.force_login(self.user)

        self.high = StudataSchool.objects.create(
            dw_schoolkey=11, state_id="HS-1", short_name="High", school_name="High School",
            school_group="High School",
        )
        self.elementary = StudataSchool.objects.create(
            dw_schoolkey=12, state_id="EL-1", short_name="Elem", school_name="Elementary School",
            school_group="Elementary School",
        )
        self.students = {
            "HS-1": CurrentStudent.objects.create(
                stu_annual_key="period-hs", ssid="period-hs", student_id="period-hs",
                student_name="High Student", grade="09",
                school_state_id="HS-1", school_year="2025-2026",
            ),
            "EL-1": CurrentStudent.objects.create(
                stu_annual_key="period-el", ssid="period-el", student_id="period-el",
                student_name="Elem Student", grade="03",
                school_state_id="EL-1", school_year="2025-2026",
            ),
        }

    def _get(self, school, teacher="Smith", period=None):
        student = self.students[school]
        url = reverse("desi:assign_accommodations", args=[student.stu_annual_key])
        params = {"school": school, "teacher": teacher}
        if period is not None:
            params["period"] = period
        return self.client.get(url, params)

    def test_period_sort_and_elementary_helpers(self):
        from desi.views import _is_elementary_school, _period_sort_key
        self.assertTrue(_is_elementary_school(self.elementary))
        self.assertFalse(_is_elementary_school(self.high))
        self.assertFalse(_is_elementary_school(None))
        self.assertEqual(
            sorted(["10", "2", "1", "A"], key=_period_sort_key),
            ["1", "2", "10", "A"],
        )

    @mock.patch("desi.views._schedule_periods", return_value=["1", "2", "10"])
    def test_period_filter_shows_for_non_elementary_with_teacher(self, mock_periods):
        response = self._get("HS-1")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["show_period_filter"])
        self.assertEqual(response.context["periods"], ["1", "2", "10"])
        content = response.content.decode()
        self.assertIn('id="period"', content)
        self.assertIn("Period 10", content)

    @mock.patch("desi.views._schedule_periods", return_value=["1"])
    def test_period_filter_hidden_for_elementary(self, mock_periods):
        response = self._get("EL-1")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["show_period_filter"])
        self.assertNotIn('id="period"', response.content.decode())
        mock_periods.assert_not_called()

    def test_period_filter_hidden_without_teacher(self):
        response = self._get("HS-1", teacher="")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["show_period_filter"])

    @mock.patch("desi.views._schedule_student_keys", return_value=[])
    @mock.patch("desi.views._schedule_periods", return_value=["1", "2"])
    def test_period_filter_narrows_using_schedule(self, mock_periods, mock_keys):
        response = self._get("HS-1", period="2")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["period_filter"], "2")
        mock_keys.assert_called_once()
        args = mock_keys.call_args.args
        self.assertEqual(args[1], "HS-1")
        self.assertEqual(args[2], "2")
        self.assertEqual(args[3], "Smith")


class DesiSupportDefaultRuleTests(TestCase):
    """Default assignment rules: Managed-by + criteria (demographic & results)."""

    def setUp(self):
        cache.delete('landing:project_config')
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="defrule@test.com",
            email="defrule@test.com",
            password="password123",
        )
        self.user.roles.add(role)
        self.client.force_login(self.user)

        self.assessment = Assessment.objects.create(name="Default Rule Math", subject="Math", is_active=True)
        self.acc = Accommodation.objects.create(name="Calculator", code="CALC", managed_by="district")
        self.opt = AccommodationOption.objects.create(accommodation=self.acc, name="Scientific", code="SCI")
        self.assessment.supported_options.add(self.opt)

        self.student = CurrentStudent.objects.create(
            stu_annual_key="dr-student", ssid="dr-ssid", student_id="dr",
            student_name="Default Rule Student", grade="05", school_year="2025-2026",
        )

    def _global_rule(self):
        rule = SupportDefaultRule.objects.create(
            name="Global Calc", support=self.acc, scope="all", is_active=True
        )
        rule.options.add(self.opt)
        return rule

    def test_global_rule_locks_option_for_every_student(self):
        self._global_rule()
        locks = default_support_locks(self.student)
        self.assertIn(self.opt.id, locks["option_ids"])
        self.assertIn(self.acc.id, locks["support_ids"])
        self.assertEqual(locks["managers"][self.acc.id], "District")

    def test_demographic_criteria_rule(self):
        rule = SupportDefaultRule.objects.create(
            name="Grade 5", support=self.acc, scope="criteria", is_active=True
        )
        rule.options.add(self.opt)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule, criterion_type="demographic", field_name="grade",
            operator="is", value="05",
        )
        self.assertIn(self.opt.id, default_support_locks(self.student)["option_ids"])

        other = CurrentStudent.objects.create(
            stu_annual_key="dr-2", ssid="dr-2", student_id="dr2",
            student_name="Other", grade="06", school_year="2025-2026",
        )
        self.assertNotIn(self.opt.id, default_support_locks(other)["option_ids"])

    def test_result_criteria_rule(self):
        Test.objects.create(
            test_group_key="elpac-sum-2025", test_name="ELPAC",
            subject="summative", school_year="2025-2026",
        )
        Result.objects.create(
            test_id="r1", test_group_key="elpac-sum-2025", ssid=self.student.ssid,
            local_student_id="dr", school_year="2025-2026", pl="2",
        )
        rule = SupportDefaultRule.objects.create(
            name="ELPAC PL12", support=self.acc, scope="criteria", is_active=True
        )
        rule.options.add(self.opt)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule, criterion_type="result", test_name="ELPAC", test_subject="summative",
            school_year="2025-2026", result_field="pl", result_operator="in", result_value="1,2",
        )
        self.assertIn(self.opt.id, default_support_locks(self.student)["option_ids"])

        other = CurrentStudent.objects.create(
            stu_annual_key="dr-3", ssid="dr-3", student_id="dr3",
            student_name="PL3", grade="05", school_year="2025-2026",
        )
        Result.objects.create(
            test_id="r2", test_group_key="elpac-sum-2025", ssid="dr-3",
            local_student_id="dr3", school_year="2025-2026", pl="3",
        )
        self.assertNotIn(self.opt.id, default_support_locks(other)["option_ids"])

    def test_workspace_shows_locked_default_option(self):
        self._global_rule()
        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Managed by District", content)
        self.assertIn(f'<option value="{self.opt.id}" selected>Scientific</option>', content)

    def test_admin_creates_rule_with_demographic_and_result_criteria(self):
        response = self.client.post(reverse("desi:support_default_rule_create"), {
            "name": "New Default",
            "support_id": str(self.acc.id),
            "scope": "criteria",
            "is_active": "on",
            "options": [str(self.opt.id)],
            "crit_type[]": ["demographic", "result"],
            "crit_logic_op[]": ["AND", "OR"],
            "crit_field[]": ["grade", ""],
            "crit_operator[]": ["is", ""],
            "crit_value[]": ["05", ""],
            "crit_test_name[]": ["", "ELPAC"],
            "crit_test_subject[]": ["", "summative"],
            "crit_test_window[]": ["", ""],
            "crit_school_year[]": ["", "2025-2026"],
            "crit_result_field[]": ["pl", "pl"],
            "crit_result_operator[]": ["in", "in"],
            "crit_result_value[]": ["", "1,2"],
        })
        self.assertEqual(response.status_code, 302)
        rule = SupportDefaultRule.objects.get(name="New Default")
        self.assertEqual(rule.options.count(), 1)
        crits = list(rule.criteria.all())
        self.assertEqual(len(crits), 2)
        self.assertEqual(crits[0].criterion_type, "demographic")
        self.assertEqual(crits[1].criterion_type, "result")

    def test_edit_support_saves_managed_by(self):
        edit_url = reverse("desi:edit_accommodation", args=[self.acc.id])
        response = self.client.post(edit_url, {
            "name": self.acc.name,
            "code": self.acc.code,
            "category": self.acc.category,
            "managed_by": "testing_coordinator",
            "description": "",
            "order": "0",
            "option_id[]": [str(self.opt.id)],
            "option_name[]": [self.opt.name],
            "option_code[]": [self.opt.code],
            "supported_assessments": [str(self.assessment.id)],
        })
        self.assertEqual(response.status_code, 302)
        self.acc.refresh_from_db()
        self.assertEqual(self.acc.managed_by, "testing_coordinator")

    def test_criteria_popover_partial_preselects_support(self):
        response = self.client.get(
            reverse("desi:support_default_rule_create"),
            {"support": str(self.acc.id), "embed": "1"},
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        # Rendered as a bare partial (no full page chrome).
        self.assertNotIn("<html", content.lower())
        self.assertIn(f'name="support_id" value="{self.acc.id}"', content)
        self.assertIn("Options to Assign", content)

    def test_admin_dashboard_has_criteria_popover(self):
        response = self.client.get(reverse("desi:admin_dashboard"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("desi-default-rule-modal", content)
        self.assertIn("/desi/default-rules/manage/", content)

    def test_manager_lists_rules_and_allows_edit(self):
        self._global_rule()
        support_url = reverse("desi:support_default_rule_manage", args=[self.acc.id])
        response = self.client.get(support_url)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Option criteria", content)
        self.assertIn("Global Calc", content)
        self.assertIn(reverse("desi:support_default_rule_edit", args=[
            SupportDefaultRule.objects.get(name="Global Calc").id,
        ]) + "?embed=1", content)

    def test_embedded_save_returns_refreshed_manager(self):
        response = self.client.post(
            reverse("desi:support_default_rule_create") + "?embed=1",
            {
                "name": "Popover Rule",
                "support_id": str(self.acc.id),
                "scope": "all",
                "is_active": "on",
                "options": [str(self.opt.id)],
            },
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Option criteria", content)
        self.assertIn("Popover Rule", content)
        self.assertTrue(SupportDefaultRule.objects.filter(name="Popover Rule").exists())

    def test_embedded_delete_returns_refreshed_manager(self):
        rule = self._global_rule()
        response = self.client.get(
            reverse("desi:support_default_rule_delete", args=[rule.id]) + "?embed=1"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("Option criteria", response.content.decode())
        self.assertFalse(SupportDefaultRule.objects.filter(id=rule.id).exists())

    def test_result_criterion_created_from_assessment_dropdown(self):
        response = self.client.post(reverse("desi:support_default_rule_create"), {
            "name": "ELPAC Dropdown Rule",
            "support_id": str(self.acc.id),
            "scope": "criteria",
            "is_active": "on",
            "options": [str(self.opt.id)],
            "crit_type[]": ["result"],
            "crit_logic_op[]": ["AND"],
            "crit_field[]": [""], "crit_operator[]": [""], "crit_value[]": [""],
            "crit_assessment[]": ["ELPAC_summative|"],
            "crit_school_year[]": ["2025-2026"],
            "crit_result_field[]": ["pl"],
            "crit_result_operator[]": ["in"],
            "crit_result_value[]": ["1,2"],
            "crit_school[]": [""], "crit_staff[]": [""], "crit_period[]": [""],
        })
        self.assertEqual(response.status_code, 302)
        crit = SupportDefaultRuleCriteria.objects.get(rule__name="ELPAC Dropdown Rule")
        self.assertEqual(crit.criterion_type, "result")
        self.assertEqual(crit.test_name, "ELPAC_summative")
        self.assertEqual(crit.test_subject, "")
        self.assertEqual(crit.school_year, "2025-2026")
        self.assertEqual(crit.result_field, "pl")
        self.assertEqual(crit.result_value, "1,2")

    def test_schedule_criterion_created_from_school_class_dropdowns(self):
        response = self.client.post(reverse("desi:support_default_rule_create"), {
            "name": "Class Rule",
            "support_id": str(self.acc.id),
            "scope": "criteria",
            "is_active": "on",
            "options": [str(self.opt.id)],
            "crit_type[]": ["schedule"],
            "crit_logic_op[]": ["AND"],
            "crit_field[]": [""], "crit_operator[]": [""], "crit_value[]": [""],
            "crit_assessment[]": [""], "crit_school_year[]": [""],
            "crit_result_field[]": ["pl"], "crit_result_operator[]": ["in"], "crit_result_value[]": [""],
            "crit_school[]": ["SCH-X"], "crit_staff[]": ["Smith, Jane"], "crit_period[]": ["3"],
        })
        self.assertEqual(response.status_code, 302)
        crit = SupportDefaultRuleCriteria.objects.get(rule__name="Class Rule")
        self.assertEqual(crit.criterion_type, "schedule")
        self.assertEqual(crit.school_state_id, "SCH-X")
        self.assertEqual(crit.staff_name, "Smith, Jane")
        self.assertEqual(crit.period, "3")

    def test_schedule_criterion_school_level_matches_by_school(self):
        rule = SupportDefaultRule.objects.create(
            name="School Scope", support=self.acc, scope="criteria", is_active=True
        )
        rule.options.add(self.opt)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule, criterion_type="schedule", school_state_id="SCH-X",
        )
        self.student.school_state_id = "SCH-X"
        self.student.save(update_fields=["school_state_id"])
        self.assertIn(self.opt.id, default_support_locks(self.student)["option_ids"])

        self.student.school_state_id = "SCH-Y"
        self.student.save(update_fields=["school_state_id"])
        self.assertNotIn(self.opt.id, default_support_locks(self.student)["option_ids"])

    def test_schedule_criterion_teacher_and_period_uses_schedule(self):
        rule = SupportDefaultRule.objects.create(
            name="Class Scope", support=self.acc, scope="criteria", is_active=True
        )
        rule.options.add(self.opt)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule, criterion_type="schedule",
            staff_name="Smith, Jane", period="3",
        )
        fake = mock.MagicMock()
        fake.exists.return_value = True
        fake.filter.return_value = fake
        with mock.patch("desi.views.Schedule") as mock_schedule:
            mock_schedule.objects.filter.return_value = fake
            self.assertIn(self.opt.id, default_support_locks(self.student)["option_ids"])

        fake.exists.return_value = False
        with mock.patch("desi.views.Schedule") as mock_schedule:
            mock_schedule.objects.filter.return_value = fake
            self.assertNotIn(self.opt.id, default_support_locks(self.student)["option_ids"])

    def test_multiple_criteria_rules_first_match_then_default(self):
        # Distinct options per rule.
        opt_a = AccommodationOption.objects.create(accommodation=self.acc, name="Option A", code="A")
        opt_b = AccommodationOption.objects.create(accommodation=self.acc, name="Option B", code="B")
        opt_def = AccommodationOption.objects.create(accommodation=self.acc, name="Option D", code="D")

        rule_a = SupportDefaultRule.objects.create(
            name="Grade 5 -> A", support=self.acc, scope="criteria", priority=10, is_active=True)
        rule_a.options.add(opt_a)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule_a, criterion_type="demographic", field_name="grade", operator="is", value="05")

        rule_b = SupportDefaultRule.objects.create(
            name="Grade 6 -> B", support=self.acc, scope="criteria", priority=20, is_active=True)
        rule_b.options.add(opt_b)
        SupportDefaultRuleCriteria.objects.create(
            rule=rule_b, criterion_type="demographic", field_name="grade", operator="is", value="06")

        default_rule = SupportDefaultRule.objects.create(
            name="Default -> D", support=self.acc, scope="all", priority=100, is_active=True)
        default_rule.options.add(opt_def)

        # Grade 05 matches rule A (highest priority criteria).
        locks = default_support_locks(self.student)
        self.assertIn(opt_a.id, locks["option_ids"])
        self.assertNotIn(opt_b.id, locks["option_ids"])
        self.assertNotIn(opt_def.id, locks["option_ids"])

        # Grade 06 matches rule B.
        self.student.grade = "06"
        self.student.save(update_fields=["grade"])
        locks = default_support_locks(self.student)
        self.assertIn(opt_b.id, locks["option_ids"])
        self.assertNotIn(opt_a.id, locks["option_ids"])
        self.assertNotIn(opt_def.id, locks["option_ids"])

        # Grade 07 matches nothing -> default.
        self.student.grade = "07"
        self.student.save(update_fields=["grade"])
        locks = default_support_locks(self.student)
        self.assertIn(opt_def.id, locks["option_ids"])
        self.assertNotIn(opt_a.id, locks["option_ids"])
        self.assertNotIn(opt_b.id, locks["option_ids"])

    def test_rule_form_saves_priority(self):
        response = self.client.post(reverse("desi:support_default_rule_create"), {
            "name": "Prioritized Rule",
            "support_id": str(self.acc.id),
            "scope": "all",
            "priority": "5",
            "is_active": "on",
            "options": [str(self.opt.id)],
        })
        self.assertEqual(response.status_code, 302)
        rule = SupportDefaultRule.objects.get(name="Prioritized Rule")
        self.assertEqual(rule.priority, 5)


class Desi504AdminRoleTests(TestCase):
    """504 students are editable only by the 504_Admin role (school-scoped)."""

    def setUp(self):
        cache.delete('landing:project_config')
        self.project_role = Role.objects.create(name="Project Admin")
        self.school_admin_role = Role.objects.create(name="School Admin")
        self.teacher_role = Role.objects.create(name="Teacher")
        self.admin504_role, _ = Role.objects.get_or_create(
            name="504_Admin", defaults={"description": "504 editor"}
        )

        def make_user(username, role, schools):
            user = User.objects.create_user(
                username=username, email=username, password="password123"
            )
            user.roles.add(role)
            user.schools = schools
            user.save()
            return user

        self.project_admin = make_user("proj@test.com", self.project_role, [])
        self.school_admin = make_user("school@test.com", self.school_admin_role, ["SCH-504"])
        self.teacher = make_user("teacher@test.com", self.teacher_role, ["SCH-504"])
        self.admin504 = make_user("admin504@test.com", self.admin504_role, ["SCH-504"])
        self.admin504_other = make_user("admin504other@test.com", self.admin504_role, ["SCH-OTHER"])

        self.student_504 = CurrentStudent.objects.create(
            stu_annual_key="s504", ssid="s504", student_id="s504",
            student_name="Student 504", grade="04", in_504="Y",
            school_state_id="SCH-504", school_year="2025-2026",
        )
        self.student_ok = CurrentStudent.objects.create(
            stu_annual_key="sok", ssid="sok", student_id="sok",
            student_name="Student OK", grade="04", in_504="N",
            school_state_id="SCH-504", school_year="2025-2026",
        )

    def test_role_exists(self):
        self.assertTrue(Role.objects.filter(name="504_Admin").exists())

    def test_504_student_editable_only_by_504_admin(self):
        self.assertFalse(can_edit_desi(self.school_admin, student=self.student_504))
        self.assertFalse(can_edit_desi(self.teacher, student=self.student_504))
        self.assertTrue(can_edit_desi(self.admin504, student=self.student_504))
        # Project Admin has full rights to all.
        self.assertTrue(can_edit_desi(self.project_admin, student=self.student_504))

    def test_504_admin_scoped_to_assigned_school(self):
        self.assertFalse(can_edit_desi(self.admin504_other, student=self.student_504))

    def test_non_504_student_editing_unaffected(self):
        self.assertTrue(can_edit_desi(self.school_admin, student=self.student_ok))
        self.assertTrue(can_edit_desi(self.teacher, student=self.student_ok))

    def test_workspace_view_only_for_school_admin_on_504(self):
        self.client.force_login(self.school_admin)
        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student_504.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["can_edit"])

    def test_workspace_editable_for_504_admin(self):
        self.client.force_login(self.admin504)
        response = self.client.get(
            reverse("desi:assign_accommodations", args=[self.student_504.stu_annual_key])
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["can_edit"])

    def test_superuser_has_full_rights_on_504(self):
        su = User.objects.create_superuser(
            username="su504@test.com", email="su504@test.com", password="password123"
        )
        # Superusers have full rights to all students, including 504.
        self.assertTrue(can_edit_desi(su, student=self.student_504))


class DesiStateUploadExportTests(TestCase):
    """CAASPP/ELPAC TOMS upload workbooks built from DESI + SEIS data."""

    def setUp(self):
        import io  # noqa: F401
        cache.delete('landing:project_config')
        role = Role.objects.create(name="Project Admin")
        self.user = User.objects.create_user(
            username="stateupload@test.com",
            email="stateupload@test.com",
            password="password123",
        )
        self.user.roles.add(role)

        caaspp = Accommodation.objects.create(
            name="Abacus", code="B", column_letter="B", tag="CAASPP"
        )
        self.caaspp_opt = AccommodationOption.objects.create(
            accommodation=caaspp, name="Abacus is selected.", code="NEA_Abacus"
        )
        elpac = Accommodation.objects.create(
            name="Alternate Response", code="B", column_letter="B", tag="ELPAC"
        )
        self.elpac_opt = AccommodationOption.objects.create(
            accommodation=elpac, name="Alternate Response is selected.", code="NEA_AR"
        )

        self.student = CurrentStudent.objects.create(
            stu_annual_key="state-stu", ssid="1234567890", student_id="stu1",
            student_name="State Student", grade="05", school_year="2025-2026",
        )
        record, _ = StudentAccommodation.objects.get_or_create(
            student_annual_key=self.student.stu_annual_key
        )
        record.options.add(self.caaspp_opt)

    def _load(self, payload):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(payload), data_only=True)

    def test_workbook_matches_template_headers_and_places_codes(self):
        from desi.state_upload_export import build_state_upload_workbook
        from desi.state_upload_templates import CAASPP_HEADERS, CAASPP_SHEET_NAME

        wb = build_state_upload_workbook("CAASPP", [self.student])
        ws = wb.active
        self.assertEqual(ws.title, CAASPP_SHEET_NAME)
        self.assertEqual([c.value for c in ws[1]], CAASPP_HEADERS)
        row = [c.value for c in ws[2]]
        self.assertEqual(row[0], "1234567890")
        # Column B (index 1) = Abacus -> DESI-selected option code.
        self.assertEqual(row[1], "NEA_Abacus")

    def test_seis_codes_are_merged(self):
        ensure_desi_seis_table()
        DesiSeis.objects.create(
            id="seis-state", ssid=self.student.ssid, test_group="elpac",
            settings={"TTS": "NEA_AR"}, assessments=["ELPAC_summative"],
        )
        from desi.state_upload_export import build_state_upload_workbook
        from desi.state_upload_templates import ELPAC_SHEET_NAME

        wb = build_state_upload_workbook("ELPAC", [self.student])
        ws = wb.active
        self.assertEqual(ws.title, ELPAC_SHEET_NAME)
        row = [c.value for c in ws[2]]
        # ELPAC column B (index 1) = Alternate Response -> SEIS code.
        self.assertEqual(row[1], "NEA_AR")

    def test_state_uploads_filter_students_by_program_grade(self):
        from desi.state_upload_export import build_state_upload_workbook

        grades = ['K', '01', '3', '08', '9', '10', '11', '12', 'TK', 'PS']
        students = [
            CurrentStudent.objects.create(
                stu_annual_key=f"grade-{index}",
                ssid=f"90000000{index:02d}",
                grade=grade,
                school_year="2025-2026",
            )
            for index, grade in enumerate(grades, start=1)
        ]

        def exported_ssids(tag, sheet_name):
            workbook = build_state_upload_workbook(tag, students)
            sheet = workbook[sheet_name]
            return [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)
                    if sheet.cell(row=row, column=1).value]

        self.assertEqual(
            exported_ssids("CAASPP", "CAASPP Test Settings Template"),
            ["9000000003", "9000000004", "9000000007"],
        )
        self.assertEqual(
            exported_ssids("ELPAC", "ELPAC Test Settings Template"),
            ["9000000001", "9000000002", "9000000003", "9000000004",
             "9000000005", "9000000006", "9000000007", "9000000008"],
        )

    def test_endpoint_returns_zip_with_two_workbooks(self):
        import zipfile
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("desi:export_state_upload"), {"school_year": "2025-2026"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        archive = zipfile.ZipFile(io.BytesIO(response.content))
        self.assertIn("CAASPP_Upload_2025-2026.xlsx", archive.namelist())
        self.assertIn("ELPAC_Upload_2025-2026.xlsx", archive.namelist())

    def test_endpoint_returns_single_workbook(self):
        self.client.force_login(self.user)
        response = self.client.get(
            reverse("desi:export_state_upload_caaspp"), {"school_year": "2025-2026"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        import zipfile
        package = zipfile.ZipFile(io.BytesIO(response.content))
        template_xml = package.read("xl/worksheets/sheet4.xml")
        self.assertIn(b"xmlns:x14=", template_xml)
        self.assertIn(b'r="A2"', template_xml)
        wb = self._load(response.content)
        self.assertEqual(wb.active.title, "CAASPP Test Settings Template")

        response = self.client.get(
            reverse("desi:export_state_upload_elpac"), {"school_year": "2025-2026"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        wb = self._load(response.content)
        self.assertEqual(wb.active.title, "ELPAC Test Settings Template")

    def test_state_upload_accessible_with_token_without_login(self):
        token = signing.dumps(
            {"uid": self.user.id, "purpose": "desi_export"},
            salt=EXPORT_ENDPOINT_SIGNING_SALT,
        )
        # Not logged in: the signed token alone authorizes the download.
        response = self.client.get(
            reverse("desi:export_state_upload_caaspp"),
            {"auth": token, "school_year": "2025-2026"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])

        response = self.client.get(
            reverse("desi:export_state_upload"),
            {"auth": token, "school_year": "2025-2026"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")

    def test_state_upload_requires_auth(self):
        response = self.client.get(reverse("desi:export_state_upload_caaspp"))
        self.assertEqual(response.status_code, 403)

    def test_export_page_publishes_state_upload_endpoints(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("desi:export_data"))
        content = response.content.decode()
        self.assertIn("data-endpoint-url", content)
        self.assertIn("/desi/admin-dashboard/export/state-upload/caaspp/", content)
        self.assertIn("/desi/admin-dashboard/export/state-upload/elpac/", content)
